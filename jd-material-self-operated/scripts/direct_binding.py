from __future__ import annotations
import protected_core as _protected_core
import argparse
import base64
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import closing, contextmanager, nullcontext
from dataclasses import asdict, dataclass, replace
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import urlparse
import httpx
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps
import jd_material_agent as core
from jd_material_agent import PreparedSource, SpuJob, WebcliTransportError, _auto_maintain_plan_hash, _bind_state, _has_material_work, _jd_image_url, _jfs_path, _load_state, _save_state, cell_text, validate_selling_points, weighted_length

@dataclass(frozen=True)
class MaterialBindingPlan:
    prepared: PreparedSource
    output_dir: Path
    requests: tuple[dict[str, Any], ...]
    authorization_requests: tuple[dict[str, Any], ...]
    target_spus: tuple[dict[str, Any], ...]
    preflight_failures: tuple[dict[str, str], ...] = ()

@dataclass(frozen=True)
class MaterialBindingResult:
    submitted_types: int
    approved_spus: int
    pending_spus: int
    failed_spus: int
    browser_required: bool
    confirmation_required: bool
    confirm_token: str | None
    success_workbook: Path | None
    pending_workbook: Path | None
    failure_workbook: Path | None
    failures: tuple[str, ...] = ()
    failed_spu_ids: tuple[str, ...] = ()
    readback_path: Path | None = None
    readback_sha256: str = ''
    readback_deferred: bool = False

class ShortTitleBatchIncomplete(RuntimeError):
    pass

def _material_value(material: dict[str, Any]) -> str:
    material_type = int(material.get('materialType') or 0)
    if material_type == 1002:
        contents = (material.get('textMaterial') or {}).get('contents') or []
        return cell_text(contents[0]) if contents else ''
    images = (material.get('imageMaterial') or {}).get('skuImages') or []
    return _jd_image_url(images[0].get('imgUrl')) if images and images[0].get('imgUrl') else ''

def _binding_material(material_type: int, order: int, value: str, sku_id: str) -> dict[str, Any]:
    material = {'skuId': int(sku_id), 'materialType': int(material_type), 'order': int(order)}
    if material_type == 1002:
        material['textMaterial'] = {'contents': [value]}
    else:
        material['imageMaterial'] = {'imgType': 1, 'replaceType': 1, 'skuImages': [{'index': 1, 'imgUrl': _jfs_path(value)}]}
    return material

def build_material_binding_requests(job: SpuJob, state_item: dict[str, Any]) -> list[dict[str, Any]]:
    return _protected_core.call('direct_binding.build_material_binding_requests', locals())

def rejected_material_matches(request, materials_by_sku):
    matches = []
    material_type = int(request['materialType'])
    for desired in request['body']['skuMaterials']:
        sku_id = cell_text(desired['skuId'])
        targets = {(int(item['order']), _material_value(item) if material_type == 1002 else _jfs_path(_material_value(item))) for item in desired['materials']}
        for actual in materials_by_sku.get(sku_id, []):
            if int(actual.get('materialType') or 0) != material_type or int(actual.get('status') or 0) not in (2, 5):
                continue
            value = _material_value(actual)
            value = value if material_type == 1002 else _jfs_path(value)
            if (int(actual.get('order') or 0), value) in targets:
                matches.append({'skuId': sku_id, 'reason': cell_text(actual.get('reason'))})
    return matches

def _evaluate_binding_request(request: dict[str, Any], materials_by_sku: dict[str, list[dict[str, Any]]]) -> tuple[str, str]:
    pending = False
    material_type = int(request['materialType'])
    expected_by_sku = {cell_text(item['skuId']): item['materials'] for item in request['body']['skuMaterials']}
    for sku_id in request['skuIds']:
        expected = expected_by_sku[cell_text(sku_id)]
        actual = [item for item in materials_by_sku.get(cell_text(sku_id), []) if int(item.get('materialType') or 0) == material_type]
        if len(actual) != len(expected):
            return ('mismatch', f'SKU {sku_id} material count mismatch')
        for item in actual:
            if material_type == 1002:
                contents = (item.get('textMaterial') or {}).get('contents') or []
                if len(contents) != 1:
                    return ('mismatch', f'SKU {sku_id} selling-point contents mismatch')
            else:
                images = (item.get('imageMaterial') or {}).get('skuImages') or []
                if len(images) != 1 or int(images[0].get('index') or 0) != 1:
                    return ('mismatch', f'SKU {sku_id} image structure mismatch')

        def comparable(item: dict[str, Any]) -> str:
            value = _material_value(item)
            return value if material_type == 1002 else _jfs_path(value)
        expected_values = {int(item['order']): comparable(item) for item in expected}
        actual_values = {int(item.get('order') or 0): comparable(item) for item in actual}
        if expected_values != actual_values:
            return ('mismatch', f'SKU {sku_id} material order or value mismatch')
        statuses = [int(item.get('status') or 0) for item in actual]
        if any((status not in (3, 4) for status in statuses)):
            details = list(dict.fromkeys((cell_text(item.get('reason')) for item in actual if int(item.get('status') or 0) not in (3, 4) and cell_text(item.get('reason')))))
            return ('failed', f'SKU {sku_id} material was rejected' + (': ' + '; '.join(details) if details else ''))
        pending = pending or any((status == 3 for status in statuses))
    return ('pending', 'waiting for audit') if pending else ('approved', 'approved')

def build_self_operated_short_title_requests(before: list[dict[str, Any]], desired: dict[str, str], *, target_rows: int=50) -> list[dict[str, Any]]:
    return _protected_core.call('direct_binding.build_self_operated_short_title_requests', locals())

def _exact_short_title_outcomes(body, error):
    try:
        outcomes = json.JSONDecoder().raw_decode(str(error).split('outcomes=', 1)[1])[0]
    except (ValueError, IndexError, TypeError):
        return None
    expected = {row['skuId'] for row in body['reqList']}
    if not isinstance(outcomes, list) or len(outcomes) != len(body['reqList']):
        return None
    if any((not isinstance(item, dict) or item.get('successKnown') is not True or type(item.get('success')) is not bool for item in outcomes)):
        return None
    identifiers = [cell_text(item.get('skuId')) for item in outcomes]
    if set(identifiers) != expected or len(set(identifiers)) != len(identifiers):
        return None
    return outcomes

def _record_short_title_outcomes(body, error, desired, accepted, unconfirmed, rejected):
    outcomes = _exact_short_title_outcomes(body, error)
    if outcomes is None:
        return 0
    requested = {row['skuId']: row['shortTitle'] for row in body['reqList']}
    if any((desired.get(sku_id) != title for sku_id, title in requested.items())):
        return 0
    submitted = 0
    for item in outcomes:
        sku_id = cell_text(item['skuId'])
        if unconfirmed.get(sku_id) != requested[sku_id]:
            continue
        unconfirmed.pop(sku_id)
        if item['success']:
            accepted[sku_id] = requested[sku_id]
            rejected.pop(sku_id, None)
            submitted += 1
        else:
            rejected[sku_id] = {'title': requested[sku_id], 'reason': cell_text(item.get('message')), 'code': cell_text(item.get('code'))}
    return submitted

def submit_self_operated_short_titles(client, spu_ids, desired, output: Path, *, expected_before=None, defer_conflicts=False) -> dict[str, Any]:
    import upload_quarantine
    upload_quarantine.assert_write_targets(output, spu_ids)
    if getattr(client, 'short_title_transport', 'product') == 'osw':
        import osw_title_tasks
        return osw_title_tasks.submit(client, spu_ids, desired, output, expected_before=expected_before, defer_conflicts=defer_conflicts)
    if not desired:
        return {'submitted': 0, 'verified': 0, 'reused': 0, 'failed': [], 'pending': []}
    before = client.short_title_rows(spu_ids)
    target_rows = getattr(client, 'short_title_target_rows', 50)
    requests = build_self_operated_short_title_requests(before, desired, target_rows=target_rows if type(target_rows) is int else 50)
    unavailable = {cell_text(row['skuId']) for row in before if row.get('shortTitleKnown') is not True and cell_text(row['skuId']) in desired}
    ready = {sku_id: title for sku_id, title in desired.items() if sku_id not in unavailable}
    conflicting = {}
    if expected_before is not None:
        for row in before:
            sku_id = cell_text(row['skuId'])
            if sku_id in ready and cell_text(row.get('shortTitle')) not in {expected_before.get(sku_id), desired[sku_id]}:
                if not defer_conflicts:
                    raise ValueError('short title changed since the plan; conflict requires a new confirmation')
                conflicting[sku_id] = {'original': expected_before.get(sku_id), 'current': cell_text(row.get('shortTitle')), 'desired': desired[sku_id]}
    audit_path = output / '.state' / 'short-title-write.json'
    previous = json.loads(audit_path.read_text(encoding='utf-8')) if audit_path.exists() else {}
    accepted = {sku_id: title for sku_id, title in previous.get('submittedTitles', {}).items() if desired.get(sku_id) == title}
    unconfirmed = dict(previous.get('unconfirmedTitles', {}))
    rejected = dict(previous.get('rejectedTitles', {}))
    if any((not isinstance(item, dict) or (sku_id in desired and item.get('title') != desired[sku_id]) for sku_id, item in rejected.items())):
        raise ValueError('rejected short-title payload changed; reconcile before replacing')
    if 'unconfirmedTitles' not in previous and (previous.get('errors') or 'after' not in previous):
        unconfirmed.update({row['skuId']: row['shortTitle'] for body in previous.get('requests', []) for row in body['reqList'] if row['skuId'] not in accepted})
    if any((sku_id in desired and desired[sku_id] != title for sku_id, title in unconfirmed.items())):
        raise ValueError('unconfirmed short-title write differs from desired payload; reconcile before replacing')
    for error in previous.get('errors', []):
        for body in previous.get('requests', []):
            _record_short_title_outcomes(body, error, desired, accepted, unconfirmed, rejected)
    before_values = {cell_text(row['skuId']): cell_text(row['shortTitle']) for row in before if row.get('shortTitleKnown') is True}
    for sku_id, title in list(unconfirmed.items()):
        if desired.get(sku_id) == title and before_values.get(sku_id) == title:
            accepted[sku_id] = title
            unconfirmed.pop(sku_id)
    for sku_id, item in list(rejected.items()):
        if desired.get(sku_id) == item['title'] and before_values.get(sku_id) == item['title']:
            accepted[sku_id] = item['title']
            rejected.pop(sku_id)
    requests = [{'reqList': [row for row in body['reqList'] if row['skuId'] in ready and row['skuId'] not in unconfirmed and (row['skuId'] not in rejected) and (accepted.get(row['skuId']) != row['shortTitle'])]} for body in requests]
    requests = [body for body in requests if body['reqList']]
    spu_by_sku = {cell_text(row['skuId']): cell_text(row['productId']) for row in before}
    deferred_spus = {spu_by_sku[sku_id] for sku_id in set(unconfirmed) | set(rejected) | set(conflicting) if sku_id in spu_by_sku}
    conflict_spus = {spu_by_sku[sku_id] for sku_id in conflicting}
    requests = [{'reqList': [row for row in body['reqList'] if row['productId'] not in conflict_spus]} for body in requests]
    requests = [body for body in requests if body['reqList']]
    audit = {'before': before, 'requests': requests, 'submittedTitles': accepted, 'unconfirmedTitles': unconfirmed, 'rejectedTitles': rejected, 'conflictingTitles': conflicting, 'deferredSpus': sorted(deferred_spus), 'unavailableSkus': sorted(unavailable), 'requestHash': _auto_maintain_plan_hash({'requests': requests})}
    history_path = output / '.state' / 'short-title-history' / f'{time.time_ns()}-{uuid.uuid4().hex}.json'
    audit['transport'] = 'osw' if getattr(client, 'title_audit_transport', 'product') == 'osw' else 'product'

    def save_audit():
        _save_state(history_path, audit)
        _save_state(audit_path, audit)
    save_audit()
    errors = []
    submitted = 0
    for body in requests:
        body = {'reqList': [row for row in body['reqList'] if row['productId'] not in deferred_spus]}
        if not body['reqList']:
            continue
        stop_writes = False
        unconfirmed.update({row['skuId']: row['shortTitle'] for row in body['reqList']})
        save_audit()
        try:
            client.save_short_titles(body)
            submitted += len(body['reqList'])
            accepted.update({row['skuId']: row['shortTitle'] for row in body['reqList']})
            for row in body['reqList']:
                unconfirmed.pop(row['skuId'], None)
        except _protected_core.ProtectedCoreError:
            raise
        except ShortTitleBatchIncomplete as error:
            errors.append(str(error))
            submitted += _record_short_title_outcomes(body, error, desired, accepted, unconfirmed, rejected)
            deferred_spus.update((row['productId'] for row in body['reqList']))
        except Exception as error:
            errors.append(str(error))
            stop_writes = True
        audit['errors'] = errors
        audit['deferredSpus'] = sorted(deferred_spus)
        save_audit()
        if stop_writes:
            break
    if getattr(client, 'defer_readback', False) is True and (not unconfirmed):
        failed = sorted((set(rejected) | set(conflicting) | unavailable) & set(desired))
        audit.update(errors=errors, failed=failed, pending=sorted(set(accepted) & set(desired)), readbackDeferred=True)
        save_audit()
        return {'submitted': submitted, 'verified': 0, 'reused': sum((before_values.get(sku) == title for sku, title in ready.items())), 'failed': failed, 'pending': audit['pending'], 'readbackDeferred': True}
    after = client.short_title_rows(spu_ids) if ready else before
    expected = {cell_text(row['skuId']): cell_text(row['shortTitle']) for row in before if row.get('shortTitleKnown') is True}
    expected.update(desired)
    observed = {cell_text(row['skuId']): cell_text(row['shortTitle']) for row in after if row.get('shortTitleKnown') is True}
    for sku_id, title in list(unconfirmed.items()):
        if desired.get(sku_id) == title and observed.get(sku_id) == title:
            accepted[sku_id] = title
            unconfirmed.pop(sku_id)
    for sku_id, item in list(rejected.items()):
        if desired.get(sku_id) == item['title'] and observed.get(sku_id) == item['title']:
            accepted[sku_id] = item['title']
            rejected.pop(sku_id)
    failed = {sku_id for sku_id, value in expected.items() if observed.get(sku_id) != value} | unavailable
    before_ids = {cell_text(row['skuId']) for row in before}
    after_ids = {cell_text(row['skuId']) for row in after}
    failed.update(before_ids ^ after_ids)
    if len(after_ids) != len(after):
        raise ValueError('short-title readback has duplicate SKU scope')
    before_values = {cell_text(row['skuId']): cell_text(row['shortTitle']) for row in before}
    pending = sorted((sku_id for sku_id in failed if sku_id not in unavailable and sku_id not in conflicting and (sku_id in accepted) and (sku_id in observed) and (observed[sku_id] == before_values.get(sku_id))))
    failed = sorted(set(failed) - set(pending))
    audit.update({'after': after, 'errors': errors, 'failed': failed, 'pending': pending})
    save_audit()
    return {'submitted': submitted, 'verified': sum((observed.get(sku_id) == title for sku_id, title in ready.items())), 'reused': sum((before_values.get(sku_id) == title for sku_id, title in ready.items())), 'failed': failed, 'pending': pending}

def material_binding_confirm_token(plans: Iterable[MaterialBindingPlan]) -> str:
    bodies = [{'spuId': request['spuId'], 'materialType': request['materialType'], 'body': request['body']} for plan in plans for request in plan.authorization_requests]
    encoded = json.dumps(bodies, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    return 'BIND-' + hashlib.sha256(encoded).hexdigest()[:24].upper()

def _write_binding_workbook(path: Path, title: str, jobs: Iterable[SpuJob], reasons: dict[str, str]) -> Path | None:
    selected = list(jobs)
    if not selected:
        path.unlink(missing_ok=True)
        return None
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(['SPUID', '代表SKUID', 'SKU数量', '状态说明'])
    for job in selected:
        sheet.append([job.spu_id, job.representative_sku_id, job.sku_count, reasons.get(job.spu_id, '')])
    detail = workbook.create_sheet('SKU明细')
    detail.append(['SPUID', 'SKUID', '商品名称'])
    for job in selected:
        for sku_id, product_title in job.skus:
            detail.append([job.spu_id, sku_id, product_title])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path

def build_self_operated_binding_requests(job: SpuJob, state_item: dict[str, Any]) -> list[dict[str, Any]]:
    return _protected_core.call('direct_binding.build_self_operated_binding_requests', locals())

def _query_self_operated_binding_materials(client: Any, jobs: Iterable[SpuJob]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, int], dict[str, str]]:
    selected = list(jobs)
    materials_by_sku = {sku_id: [] for job in selected for sku_id in job.sku_ids}
    sell_limits: dict[str, int] = {}
    failures: dict[str, str] = {}
    if not selected:
        return (materials_by_sku, sell_limits, failures)
    inspect = client.inspect_material_details if callable(getattr(type(client), 'inspect_material_details', None)) else client.inspect_spus
    records = inspect([job.spu_id for job in selected])
    returned_spus = [cell_text(item.get('spuId')) for item in records]
    expected_spus = [job.spu_id for job in selected]
    duplicates = {spu_id for spu_id in returned_spus if returned_spus.count(spu_id) > 1}
    for spu_id in set(expected_spus) - set(returned_spus):
        failures[spu_id] = 'self-operated material readback omitted the SPU'
    for spu_id in duplicates:
        failures[spu_id] = 'self-operated material readback duplicated the SPU'
    jobs_by_spu = {job.spu_id: job for job in selected}
    for record in records:
        spu_id = cell_text(record.get('spuId'))
        if spu_id not in jobs_by_spu or spu_id in duplicates:
            continue
        if record.get('errors'):
            failures[spu_id] = 'self-operated material readback failed'
            continue
        job = jobs_by_spu[spu_id]
        sku_rows = {cell_text(item.get('skuId')): item for item in record.get('skus') or []}
        if set(sku_rows) != set(job.sku_ids) or len(sku_rows) != len(job.sku_ids):
            failures[spu_id] = 'self-operated material readback returned an incomplete SKU set'
            continue
        text_rows = {cell_text(item.get('skuId')): item for item in record.get('textChildren') or []}
        if job.needs_selling_points and set(text_rows) != set(job.sku_ids):
            failures[spu_id] = 'self-operated selling-point readback returned an incomplete SKU set'
            continue
        for sku_id, sku in sku_rows.items():
            seen: set[tuple[int, int, int, str]] = set()
            for item in sku.get('materials') or []:
                material_type = int(item.get('type') or 0)
                order = int(item.get('order') or 0)
                status = int(item.get('status') or 0)
                url = cell_text(item.get('url'))
                key = (material_type, order, status, url)
                if not material_type or not order or key in seen:
                    continue
                seen.add(key)
                materials_by_sku[sku_id].append({'materialType': material_type, 'order': order, 'status': status, 'imageMaterial': {'skuImages': [{'index': 1, 'imgUrl': url}]}, **({'reason': cell_text(item['reason'])} if item.get('reason') else {})})
            text = text_rows.get(sku_id) or {}
            points = [cell_text(value) for value in text.get('sellPoints') or [] if cell_text(value)]
            sell_limits[sku_id] = int(text.get('sellMaxNum') or 3)
            for order, point in enumerate(points, 1):
                materials_by_sku[sku_id].append({'materialType': 1002, 'order': order, 'status': 4, 'textMaterial': {'contents': [point]}})
    return (materials_by_sku, sell_limits, failures)

def _self_operated_request_status(request: dict[str, Any], materials_by_sku: dict[str, list[dict[str, Any]]]) -> tuple[str, str]:
    return _evaluate_binding_request(request, materials_by_sku)

def _self_operated_binding_snapshot(request, materials_by_sku):
    material_type = int(request['materialType'])
    return {sku_id: sorted([[int(item.get('order') or 0), int(item.get('status') or 0), _material_value(item)] for item in materials_by_sku.get(sku_id, []) if int(item.get('materialType') or 0) == material_type]) for sku_id in request['skuIds']}

def _self_operated_receipt_status(request, materials_by_sku, receipt):
    status, reason = _self_operated_request_status(request, materials_by_sku)
    if status == 'mismatch' and receipt.get('fingerprint') == request['fingerprint'] and (receipt.get('before') == _self_operated_binding_snapshot(request, materials_by_sku)):
        return ('pending', 'accepted; readback still shows the pre-submission value')
    return (status, reason)

def _self_operated_persisted_binding_status(request, materials_by_sku, state_item):
    material_type = str(request['materialType'])
    intents = state_item.get('material_submission_intents', {})
    intent = intents.get(material_type)
    receipt = state_item.get('material_submission_receipts', {}).get(material_type, {})
    if not intent:
        return _self_operated_receipt_status(request, materials_by_sku, receipt)
    if intent.get('fingerprint') != request['fingerprint']:
        return ('failed', 'unconfirmed binding differs from the current payload; reconcile before replacing')
    status, reason = _evaluate_binding_request(request, materials_by_sku)
    if status in {'approved', 'pending'}:
        state_item.setdefault('material_submission_receipts', {})[material_type] = {'fingerprint': request['fingerprint'], 'before': intent.get('before'), 'reconciledAt': int(time.time())}
        intents.pop(material_type)
        return (status, reason)
    return ('failed', 'binding submission unconfirmed; reconcile before retrying: ' + reason)

def prepare_self_operated_binding_plan(client: Any, prepared: PreparedSource, output_dir: str | Path, *, selected_spu_ids: set[str] | None=None) -> MaterialBindingPlan:
    import upload_quarantine
    upload_quarantine.assert_write_targets(output_dir, selected_spu_ids if selected_spu_ids is not None else {job.spu_id for job in prepared.jobs})
    output = Path(output_dir)
    state_path = output / '.state' / 'task.json'
    state = _load_state(state_path)
    _bind_state(state, prepared, state_path)
    material_jobs = [job for job in prepared.jobs if _has_material_work(job) and (selected_spu_ids is None or job.spu_id in selected_spu_ids)]
    materials_by_sku, sell_limits, query_failures = _query_self_operated_binding_materials(client, material_jobs)
    _save_state(output / '.state' / f'binding-before-{time.time_ns()}.json', {'materialsBySku': materials_by_sku, 'queryFailures': query_failures})
    requests: list[dict[str, Any]] = []
    authorization_requests: list[dict[str, Any]] = []
    preflight_failures: list[dict[str, str]] = []
    for job in material_jobs:
        state_item = state['jobs'].setdefault(job.spu_id, {})
        binding_state = state_item.setdefault('material_bindings', {})
        if job.spu_id in query_failures:
            reason = query_failures[job.spu_id]
            preflight_failures.append({'spuId': job.spu_id, 'reason': reason})
            binding_state['preflight'] = {'status': 'failed', 'reason': reason, 'updatedAt': int(time.time())}
            continue
        try:
            desired = build_self_operated_binding_requests(job, state_item)
            for request in desired:
                if int(request['materialType']) == 1002:
                    point_count = len(request['wireBody']['apiWareTextMaterialInfo']['sellPoint'])
                    if any((sell_limits.get(sku_id, 0) < point_count for sku_id in job.sku_ids)):
                        raise ValueError(f'SPU {job.spu_id} selling points exceed sellMaxNum')
        except _protected_core.ProtectedCoreError:
            raise
        except Exception as error:
            preflight_failures.append({'spuId': job.spu_id, 'reason': str(error)})
            binding_state['preflight'] = {'status': 'failed', 'reason': str(error), 'updatedAt': int(time.time())}
            continue
        binding_state.pop('preflight', None)
        authorization_requests.extend(desired)
        for request in desired:
            material_type = int(request['materialType'])
            receipt = state_item.get('material_submission_receipts', {}).get(str(material_type), {})
            status, reason = _self_operated_persisted_binding_status(request, materials_by_sku, state_item)
            current = {'fingerprint': request['fingerprint'], 'status': status, 'reason': reason, 'updatedAt': int(time.time())}
            rejected = rejected_material_matches(request, materials_by_sku)
            if rejected:
                binding_state[str(material_type)] = {**current, 'status': 'failed', 'reason': 'same rejected material requires corrected replacement; not resubmitted', 'rejectedMaterials': rejected}
                continue
            if status in {'approved', 'pending'}:
                binding_state[str(material_type)] = current
                continue
            if receipt.get('fingerprint') == request['fingerprint']:
                binding_state[str(material_type)] = {**current, 'status': 'failed'}
                continue
            if str(material_type) in state_item.get('material_submission_intents', {}):
                binding_state[str(material_type)] = {**current, 'status': 'failed'}
                continue
            if material_type != 1002:
                existing = any((int(item.get('materialType') or 0) == material_type for sku_id in request['skuIds'] for item in materials_by_sku.get(cell_text(sku_id), [])))
                request['wireBody']['apiMaterialPicInfo']['existNoReplace'] = not existing
            binding_state[str(material_type)] = {**current, 'status': 'planned'}
            request['beforeSnapshot'] = _self_operated_binding_snapshot(request, materials_by_sku)
            requests.append(request)
    _save_state(state_path, state)
    targets = []
    for job in prepared.jobs:
        material_types = sorted({int(request['materialType']) for request in requests if request['spuId'] == job.spu_id})
        if material_types:
            targets.append({'spuId': job.spu_id, 'skuCount': job.sku_count, 'materialTypes': material_types, 'overwrite': True})
    return MaterialBindingPlan(prepared, output, tuple(requests), tuple(authorization_requests), tuple(targets), tuple(preflight_failures))

def _interleave_self_operated_binding_requests(requests):
    groups = {}
    for request in requests:
        groups.setdefault(request['spuId'], deque()).append(request)
    pending = deque(groups.values())
    while pending:
        group = pending.popleft()
        yield group.popleft()
        if group:
            pending.append(group)

def _self_operated_binding_request_id(request):
    return f"{request['spuId']}:{int(request['materialType'])}:{request['fingerprint']}"

def _merge_self_operated_image_requests(requests):
    groups = {}
    for request in requests:
        if int(request['materialType']) == 1002:
            key = (request['spuId'], 'text', request['fingerprint'])
        else:
            picture = request['wireBody']['apiMaterialPicInfo']
            if str(picture['apiRelativeSkus']['productId']) != str(request['spuId']) or list(picture['apiRelativeSkus']['skuIds']) != list(request['skuIds']):
                raise ValueError('mixed binding wire scope differs from the authorized request')
            key = (request['spuId'], tuple(request['skuIds']), picture['sceneType'])
        groups.setdefault(key, []).append(request)
    merged = []
    for members in groups.values():
        if len(members) == 1:
            merged.append(members[0])
            continue
        if len({item['materialType'] for item in members}) != len(members):
            merged.extend(members)
            continue
        picture = dict(members[0]['wireBody']['apiMaterialPicInfo'])
        picture['imgList'] = [image for item in members for image in item['wireBody']['apiMaterialPicInfo']['imgList']]
        picture['existNoReplace'] = all((item['wireBody']['apiMaterialPicInfo']['existNoReplace'] for item in members))
        merged.append({**members[0], 'members': members, 'wireBody': {'apiMaterialPicInfo': picture}, 'fingerprint': _auto_maintain_plan_hash({'members': [_self_operated_binding_request_id(item) for item in members]})})
    return merged

def _self_operated_binding_batch_source(requests, expected_erp, *, concurrency=1, template_only=False):
    if not expected_erp or not requests or type(concurrency) is not int or (concurrency < 1):
        raise ValueError('binding batch requires an ERP, requests and positive concurrency')
    packets = [{'requestId': _self_operated_binding_request_id(request), 'spuId': request['spuId'], 'kind': 'saveSellingPoints' if int(request['materialType']) == 1002 else 'bindImages', 'body': request['wireBody']} for request in requests]
    if len({item['spuId'] for item in packets}) != len(packets):
        raise ValueError('binding batch cannot repeat a product')
    scripts = Path(__file__).parent
    single = (scripts / 'self_operated_material_binding_bridge.js').read_text(encoding='utf-8')
    single = single.replace('__JD_SELF_OPERATED_BINDING_ACTION__', 'singleAction')
    source = (scripts / 'self_operated_material_binding_batch_bridge.js').read_text(encoding='utf-8')
    source = source.replace('__JD_SELF_OPERATED_BINDING_SINGLE_SOURCE__', single.rstrip().removesuffix(';'))
    action = {'expectedErp': expected_erp, 'requests': packets, 'concurrency': concurrency}
    if template_only:
        return (source, action)
    return source.replace('__JD_SELF_OPERATED_BINDING_BATCH_ACTION__', json.dumps(action, ensure_ascii=False))

def _self_operated_binding_batches(requests, expected_erp, *, target_size=5, cache_bridges=False):
    if type(target_size) is not int or target_size < 1:
        raise ValueError('binding target size must be a positive integer')
    group = []
    for request in requests:
        if group and (len(group) >= target_size or any((item['spuId'] == request['spuId'] for item in group)) or (not cache_bridges and len(_self_operated_binding_batch_source([*group, request], expected_erp).encode('utf-8')) >= 12000)):
            yield group
            group = []
        group.append(request)
    if group:
        yield group

def _validate_self_operated_binding_batch_reply(requests, reply):
    if not isinstance(reply, dict) or type(reply.get('stopped')) is not bool:
        raise ValueError('binding batch receipt has no explicit stopped state')
    rows = reply.get('results')
    identifiers = [_self_operated_binding_request_id(request) for request in requests]
    if not isinstance(rows, list) or len(rows) != len(identifiers) or any((not isinstance(row, dict) or row.get('status') not in {'submitted', 'unconfirmed', 'not-submitted'} for row in rows)):
        raise ValueError('binding batch receipt count or status is invalid')
    received = [row.get('requestId') for row in rows]
    if any((not isinstance(identifier, str) for identifier in received)) or set(received) != set(identifiers) or len(set(received)) != len(received):
        raise ValueError('binding batch receipt identities differ from the request')
    by_identifier = {row['requestId']: row for row in rows}
    ordered = [by_identifier[identifier] for identifier in identifiers]
    statuses = [row['status'] for row in ordered]
    if not reply['stopped'] and 'not-submitted' in statuses:
        raise ValueError('binding batch receipt has unsent requests without a stop')
    if reply['stopped']:
        if reply.get('concurrent') is True:
            if 'unconfirmed' not in statuses:
                raise ValueError('concurrent stop has no uncertain request')
            return ordered
        attempted = statuses[:statuses.index('not-submitted')] if 'not-submitted' in statuses else statuses
        if not attempted or attempted[-1] != 'unconfirmed' or any((status != 'not-submitted' for status in statuses[len(attempted):])):
            raise ValueError('binding batch stopped receipt does not have an exact unsent tail')
    return ordered

def _execute_self_operated_binding_batches(client, plan, state, state_path, request_errors):

    def is_pending(request):
        item = state['jobs'].get(request['spuId'], {})
        material_type = str(int(request['materialType']))
        receipt = item.get('material_submission_receipts', {}).get(material_type, {})
        return material_type not in item.get('material_submission_intents', {}) and receipt.get('fingerprint') != request['fingerprint']
    eligible = [request for request in plan.requests if is_pending(request)]
    if getattr(client, 'merge_image_types', False) is True:
        eligible = _merge_self_operated_image_requests(eligible)
    eligible = list(_interleave_self_operated_binding_requests(eligible))
    submitted = 0
    for requests in _self_operated_binding_batches(eligible, client.expected_erp, target_size=getattr(client, 'binding_batch_size', 5), cache_bridges=getattr(client, 'cache_bridges', False) is True):
        requests = [request for request in requests if all((is_pending(member) for member in request.get('members', [request])))]
        if not requests:
            continue
        for request in [member for packet in requests for member in packet.get('members', [packet])]:
            item = state['jobs'].setdefault(request['spuId'], {})
            item.setdefault('material_submission_intents', {})[str(int(request['materialType']))] = {'fingerprint': request['fingerprint'], 'before': request.get('beforeSnapshot'), 'startedAt': int(time.time())}
        _save_state(state_path, state)
        reply = client.bind_self_operated_batch(requests)
        outcomes = _validate_self_operated_binding_batch_reply(requests, reply)
        expanded = [(member, outcome) for packet, outcome in zip(requests, outcomes) for member in packet.get('members', [packet])]
        for request, outcome in expanded:
            item = state['jobs'][request['spuId']]
            material_type = str(int(request['materialType']))
            status = outcome['status']
            reason = cell_text(outcome.get('error')) or ('submitted' if status == 'submitted' else 'binding batch did not confirm submission')
            if status == 'submitted':
                item.setdefault('material_submission_receipts', {})[material_type] = {'fingerprint': request['fingerprint'], 'before': request.get('beforeSnapshot'), 'acceptedAt': int(time.time())}
                submitted += 1
            if status in {'submitted', 'not-submitted'}:
                item['material_submission_intents'].pop(material_type)
            if status != 'submitted':
                request_errors[request['spuId'], int(request['materialType'])] = reason
            item.setdefault('material_bindings', {})[material_type] = {'fingerprint': request['fingerprint'], 'status': 'submitted' if status == 'submitted' else 'failed', 'reason': reason, 'updatedAt': int(time.time())}
            _save_state(state_path, state)
        if reply['stopped']:
            raise RuntimeError('binding batch stopped; preserve receipts and reconcile uncertain requests before resuming')
    return submitted

def execute_self_operated_binding_plan(client: Any, plan: MaterialBindingPlan) -> MaterialBindingResult:
    import upload_quarantine
    upload_quarantine.assert_write_targets(plan.output_dir, {request['spuId'] for request in (*plan.requests, *plan.authorization_requests)})
    state_path = plan.output_dir / '.state' / 'task.json'
    state = _load_state(state_path)
    _bind_state(state, plan.prepared, state_path)
    request_errors: dict[tuple[str, int], str] = {}
    requests_by_spu: dict[str, list[dict[str, Any]]] = {}
    for request in plan.authorization_requests:
        requests_by_spu.setdefault(request['spuId'], []).append(request)
    batch_supported = callable(getattr(type(client), 'bind_self_operated_batch', None)) and bool(getattr(client, 'expected_erp', ''))
    submitted = _execute_self_operated_binding_batches(client, plan, state, state_path, request_errors) if batch_supported else 0
    for request in () if batch_supported else _interleave_self_operated_binding_requests(plan.requests):
        spu_id = request['spuId']
        material_type = int(request['materialType'])
        state_item = state['jobs'].setdefault(spu_id, {})
        binding_state = state_item.setdefault('material_bindings', {})
        intents = state_item.setdefault('material_submission_intents', {})
        receipt = state_item.get('material_submission_receipts', {}).get(str(material_type), {})
        if str(material_type) in intents or receipt.get('fingerprint') == request['fingerprint']:
            continue
        intents[str(material_type)] = {'fingerprint': request['fingerprint'], 'before': request.get('beforeSnapshot'), 'startedAt': int(time.time())}
        _save_state(state_path, state)
        try:
            if material_type == 1002:
                client.save_self_operated_selling_points(request['wireBody'])
            else:
                client.bind_self_operated_images(request['wireBody'])
            state['jobs'][spu_id].setdefault('material_submission_receipts', {})[str(material_type)] = {'fingerprint': request['fingerprint'], 'before': request.get('beforeSnapshot'), 'acceptedAt': int(time.time())}
            intents.pop(str(material_type))
            binding_state[str(material_type)] = {'fingerprint': request['fingerprint'], 'status': 'submitted', 'reason': 'submitted', 'updatedAt': int(time.time())}
            submitted += 1
        except _protected_core.ProtectedCoreError:
            raise
        except Exception as error:
            request_errors[spu_id, material_type] = str(error)
            binding_state[str(material_type)] = {'fingerprint': request['fingerprint'], 'status': 'failed', 'reason': str(error), 'updatedAt': int(time.time())}
        finally:
            _save_state(state_path, state)
    material_jobs = [job for job in plan.prepared.jobs if job.spu_id in requests_by_spu]
    if getattr(client, 'defer_readback', False) is True and (not any((item.get('material_submission_intents') for item in state['jobs'].values()))):
        readback_path = plan.output_dir / '.state' / f'binding-deferred-{time.time_ns()}.json'
        _save_state(readback_path, {'materialsBySku': {}, 'queryFailures': {}, 'readbackDeferred': True})
        failed = {item['spuId']: item['reason'] for item in plan.preflight_failures}
        return MaterialBindingResult(submitted, 0, 0, len(failed), False, False, None, None, None, None, failures=tuple(failed.values()), failed_spu_ids=tuple(failed), readback_path=readback_path, readback_sha256=hashlib.sha256(readback_path.read_bytes()).hexdigest(), readback_deferred=True)
    materials_by_sku: dict[str, list[dict[str, Any]]] = {}
    query_failures: dict[str, str] = {}
    try:
        materials_by_sku, _, query_failures = _query_self_operated_binding_materials(client, material_jobs)
    except _protected_core.ProtectedCoreError:
        raise
    except Exception as error:
        query_failures = {job.spu_id: str(error) for job in material_jobs}
    readback_path = plan.output_dir / '.state' / f'binding-after-{time.time_ns()}.json'
    _save_state(readback_path, {'materialsBySku': materials_by_sku, 'queryFailures': query_failures})
    readback_sha256 = hashlib.sha256(readback_path.read_bytes()).hexdigest()
    for request in plan.authorization_requests:
        spu_id = request['spuId']
        material_type = int(request['materialType'])
        key = (spu_id, material_type)
        binding_state = state['jobs'].setdefault(spu_id, {}).setdefault('material_bindings', {})
        if spu_id in query_failures:
            status, reason = ('failed', query_failures[spu_id])
        else:
            status, reason = _self_operated_persisted_binding_status(request, materials_by_sku, state['jobs'][spu_id])
            if status == 'mismatch':
                status = 'failed'
            if key in request_errors and status not in {'approved', 'pending'}:
                status, reason = ('failed', request_errors[key])
        binding_state[str(material_type)] = {'fingerprint': request['fingerprint'], 'status': status, 'reason': reason, 'updatedAt': int(time.time())}
    _save_state(state_path, state)
    approved_jobs: list[SpuJob] = []
    pending_jobs: list[SpuJob] = []
    failed_jobs: list[SpuJob] = []
    reasons = {item['spuId']: item['reason'] for item in plan.preflight_failures}
    for job in plan.prepared.jobs:
        planned = requests_by_spu.get(job.spu_id, [])
        if not planned and job.spu_id not in reasons:
            continue
        bindings = state['jobs'].setdefault(job.spu_id, {}).get('material_bindings') or {}
        statuses = [bindings.get(str(item['materialType']), {}) for item in planned]
        if job.spu_id in reasons or any((item.get('status') == 'failed' for item in statuses)):
            failed_jobs.append(job)
            details = [item.get('reason') for item in statuses if item.get('status') == 'failed']
            reasons[job.spu_id] = reasons.get(job.spu_id) or '；'.join(dict.fromkeys(filter(None, details))) or 'binding failed'
        elif any((item.get('status') in {'pending', 'submitted'} for item in statuses)):
            pending_jobs.append(job)
            reasons[job.spu_id] = 'waiting for audit status 4'
        elif statuses and all((item.get('status') == 'approved' for item in statuses)):
            approved_jobs.append(job)
            reasons[job.spu_id] = 'all planned materials approved'
        else:
            failed_jobs.append(job)
            reasons[job.spu_id] = 'binding state is incomplete'
    success_path = _write_binding_workbook(plan.output_dir / '素材绑定成功清单.xlsx', '绑定成功SPU', approved_jobs, reasons)
    pending_path = _write_binding_workbook(plan.output_dir / '素材绑定审核中清单.xlsx', '审核中SPU', pending_jobs, reasons)
    failure_path = _write_binding_workbook(plan.output_dir / '素材绑定失败SPU清单.xlsx', '绑定失败SPU', failed_jobs, reasons)
    failures = tuple((f'SPU {job.spu_id}: {reasons[job.spu_id]}' for job in failed_jobs))
    return MaterialBindingResult(submitted, len(approved_jobs), len(pending_jobs), len(failed_jobs), False, False, None, success_path, pending_path, failure_path, failures, tuple((job.spu_id for job in failed_jobs)), readback_path, readback_sha256)

def confirm_self_operated_writes(output: Path, payload: dict, token: str) -> str:
    payload = json.loads(json.dumps(payload, ensure_ascii=False))
    expected = 'WRITE-' + _auto_maintain_plan_hash(payload)
    path = output / '.state' / 'self-operated-writes.json'
    if token:
        if token != expected or not path.exists():
            raise ValueError('write confirmation no longer matches the exact payload')
        saved = json.loads(path.read_text(encoding='utf-8'))
        if saved != {'confirmToken': expected, 'payload': payload}:
            raise ValueError('saved write confirmation changed; prepare a new dry-run')
    else:
        _save_state(path, {'confirmToken': expected, 'payload': payload})
    return expected

class WebcliSelfOperatedClient(core.WebcliSelfOperatedClient):

    def osw_title_action(self, action):
        client = getattr(self, '_osw_title_client', None)
        if client is None:
            client = core.WebcliSelfOperatedClient(profile=self.profile, session=self.session + '-osw-titles', timeout=self.timeout)
            self._osw_title_client = client
        if not client.initialized:
            client._ensure_browser_page('https://osw.jd.com/site-fe/batch-task/create')
            frames = client._run_webcli(['frames'])
            frame = next((item['index'] for item in frames if item.get('url') == 'https://commodity-backend-fe-pro.local-pf.jd.com/batch-task/create'), None)
            if frame is None:
                raise ValueError('OSW task iframe unavailable')
            client.osw_frame = frame
            client.initialized = True
        frame = client.osw_frame
        source = Path(__file__).with_name('self_operated_osw_title_bridge.js').read_text(encoding='utf-8')
        source = source.replace('__JD_OSW_TITLE_ACTION__', json.dumps(action, ensure_ascii=False))
        if len(source.encode('utf-8')) < 12000:
            return client._run_webcli(['eval', '--frame', str(frame), source])
        key = '__jdOswSource_' + uuid.uuid4().hex
        client._run_webcli(['eval', '--frame', str(frame), f'window.{key}=[];true'])
        for offset in range(0, len(source), 3000):
            client._run_webcli(['eval', '--frame', str(frame), f'window.{key}.push({json.dumps(source[offset:offset + 3000], ensure_ascii=True)});true'])
        try:
            return client._run_webcli(['eval', '--frame', str(frame), f'(0,eval)(window.{key}.join(""))'])
        finally:
            try:
                client._run_webcli(['eval', '--frame', str(frame), f'delete window.{key}'])
            except _protected_core.ProtectedCoreError:
                raise
            except Exception:
                pass

    def _image_http(self):
        from image_space_http import ImageSpaceHttpClient
        client = getattr(self, '_http_image_client', None)
        if client is None:
            client = self._http_image_client = ImageSpaceHttpClient(self.expected_erp, self.profile)
        if (client.erp, client.profile) != (self.expected_erp, self.profile):
            raise ValueError('cached image transport identity changed')
        return client

    def close_image_transport(self):
        client = getattr(self, '_http_image_client', None)
        if client is not None:
            client.close()
            self._http_image_client = None

    def upload_images(self, paths, category_id, *, expected_hashes=None):
        if getattr(self, 'image_transport', 'browser') == 'direct-http':
            return self._image_http().upload_images(paths, category_id, expected_hashes=expected_hashes)
        return super().upload_images(paths, category_id, expected_hashes=expected_hashes)

    def image_page(self, category_id, page, page_size, query=''):
        if getattr(self, 'image_transport', 'browser') == 'direct-http':
            return self._image_http().image_page(category_id, page, page_size, query)
        return super().image_page(category_id, page, page_size, query)

    def image_pages(self, category_id, queries, page_size=50):
        if getattr(self, 'image_transport', 'browser') == 'direct-http':
            return self._image_http().image_pages(category_id, queries, page_size)
        if not self.expected_erp or not 1 <= len(queries) <= 10 or len(set(queries)) != len(queries):
            raise ValueError('image lookup requires an ERP and one to ten distinct queries')
        if not self.image_space_initialized:
            self._ensure_browser_page(self.image_space_url, image_space=True)
            self.image_space_initialized = True
        source = Path(__file__).with_name('self_operated_image_lookup_batch_bridge.js').read_text(encoding='utf-8')
        single = self.upload_bridge_path.read_text(encoding='utf-8').replace('__JD_SELF_OPERATED_UPLOAD_ACTION__', 'singleAction')
        source = source.replace('__JD_SELF_OPERATED_LOOKUP_SINGLE_SOURCE__', single)
        action = {'expectedErp': self.expected_erp, 'categoryId': str(category_id), 'queries': list(queries), 'pageSize': page_size}
        return self._eval_template(source, '__JD_SELF_OPERATED_LOOKUP_BATCH_ACTION__', action, image_space=True, readonly=True, recovery_url=self.image_space_url)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.product_bridge_path = Path(__file__).with_name('self_operated_direct_product_bridge.js')

    def bind_self_operated_images(self, body: dict[str, Any]) -> dict[str, Any]:
        data = self._eval_binding({'kind': 'bindImages', 'body': body})
        if not isinstance(data, dict) or not data.get('ok'):
            errors = data.get('errors') if isinstance(data, dict) else None
            raise RuntimeError('；'.join(errors or ['self-operated image binding failed']))
        return data

    def save_self_operated_selling_points(self, body: dict[str, Any]) -> dict[str, Any]:
        data = self._eval_binding({'kind': 'saveSellingPoints', 'body': body})
        if not isinstance(data, dict) or not data.get('ok'):
            errors = data.get('errors') if isinstance(data, dict) else None
            raise RuntimeError('；'.join(errors or ['self-operated selling-point save failed']))
        return data

    def _eval_binding(self, action: dict[str, Any]) -> Any:
        if not self.initialized:
            self._ensure_browser_page(self.page_url)
            self.initialized = True
        source = self.binding_bridge_path.read_text(encoding='utf-8')
        return self._eval_template(source, '__JD_SELF_OPERATED_BINDING_ACTION__', {**action, 'expectedErp': self.expected_erp})

    def bind_self_operated_batch(self, requests):
        source, action = _self_operated_binding_batch_source(requests, self.expected_erp, concurrency=getattr(self, 'binding_concurrency', 1), template_only=True)
        if not self.initialized:
            self._ensure_browser_page(self.page_url)
            self.initialized = True
        return self._eval_template(source, '__JD_SELF_OPERATED_BINDING_BATCH_ACTION__', action)

    def save_short_titles(self, body: dict[str, Any]) -> dict[str, Any]:
        try:
            result = self._eval_products({'kind': 'saveShortTitles', 'body': body})
            if isinstance(result, dict) and result.get('ok') is False:
                reason = cell_text(result.get('error'))
                if reason.startswith('short-title submission incomplete;'):
                    raise ShortTitleBatchIncomplete(reason)
            if not isinstance(result, dict) or result.get('ok') is not True or type(result.get('submitted')) is not int or (result['submitted'] != len(body.get('reqList', []))):
                raise RuntimeError('short-title submission did not return an exact positive acceptance')
            return result
        except _protected_core.ProtectedCoreError:
            raise
        except RuntimeError as error:
            if not isinstance(error, WebcliTransportError) and str(error).startswith('short-title submission incomplete;'):
                raise ShortTitleBatchIncomplete(str(error)) from error
            raise

def create_direct_client(args):
    base = core.create_self_operated_erp_client(args)
    client = WebcliSelfOperatedClient.__new__(WebcliSelfOperatedClient)
    client.__dict__.update(base.__dict__)
    client.image_transport = getattr(args, 'image_transport', 'browser')
    client.short_title_transport = getattr(args, 'short_title_transport', 'product')
    client.product_bridge_path = Path(__file__).with_name('self_operated_direct_product_bridge.js')
    return client

def __getattr__(name):
    return getattr(core, name)
