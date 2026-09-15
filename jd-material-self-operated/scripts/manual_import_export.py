from __future__ import annotations
import argparse
import csv
import hashlib
from io import BytesIO
import json
from pathlib import Path
import posixpath
import re
from urllib.parse import urlparse
from xml.etree import ElementTree
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import load_workbook
ASSETS = Path(__file__).resolve().parents[1] / 'assets'
MATERIAL_TEMPLATE = ASSETS / '商品素材导入模板-通用场景.xlsx'
TITLE_TEMPLATE = ASSETS / '批量维护短标题.xlsx'
IMAGE_COLUMNS = ('transparent', 'white', 'scene1', 'scene2', 'selling')
MATERIAL_MAX_BYTES = 1048576
TITLE_MAX_BYTES = 5 * 1024 * 1024 - 1

def fill_template_bytes(template, sheet_name, rows):
    with ZipFile(template) as original:
        namespace = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        sheets = ElementTree.fromstring(original.read('xl/workbook.xml'))
        selected = next((sheet for sheet in sheets.findall('main:sheets/main:sheet', namespace) if sheet.get('name') == sheet_name))
        relationship = selected.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        targets = ElementTree.fromstring(original.read('xl/_rels/workbook.xml.rels'))
        target = next((item.get('Target') for item in targets if item.get('Id') == relationship))
        part = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        xml = original.read(part).decode('utf-8')
        root = ElementTree.fromstring(xml)
        sheet_data = root.find('main:sheetData', namespace)
        if sheet_data is None or len(sheet_data) != 1 or sheet_data[0].get('r') != '1':
            raise ValueError('official import template must contain only its original header row')
        width = len(sheet_data[0])
        fragments = []
        for row_number, values in enumerate(rows, 2):
            if len(values) != width:
                raise ValueError('import row width differs from the official template')
            cells = []
            for column, value in enumerate(values, 1):
                text = '' if value is None else str(value)
                if re.search('[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f]', text):
                    raise ValueError('import value contains an invalid XML control character')
                if text:
                    cells.append(f'<c r="{chr(64 + column)}{row_number}" t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>')
                else:
                    cells.append(f'<c r="{chr(64 + column)}{row_number}"/>')
            fragments.append(f'''<row r="{row_number}">{''.join(cells)}</row>''')
        xml = xml.replace('</sheetData>', ''.join(fragments) + '</sheetData>', 1)
        xml = re.sub('(<dimension\\s+ref=)["\\\'][^"\\\']+["\\\']', lambda match: match[1] + f'"A1:{chr(64 + width)}{len(rows) + 1}"', xml, count=1)
        buffer = BytesIO()
        with ZipFile(buffer, 'w', compression=ZIP_DEFLATED) as filled:
            for info in original.infolist():
                filled.writestr(info, xml.encode('utf-8') if info.filename == part else original.read(info.filename))
        return (buffer.getvalue(), part)

def write_parts(rows, template, sheet_name, output, prefix, *, max_bytes, max_rows=None):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    files = []
    offset = 0
    while offset < len(rows):
        count = min(len(rows) - offset, max_rows or len(rows))
        data, part = fill_template_bytes(template, sheet_name, rows[offset:offset + count])
        if len(data) > max_bytes:
            low, high, best = (1, count - 1, None)
            while low <= high:
                middle = (low + high) // 2
                candidate, part = fill_template_bytes(template, sheet_name, rows[offset:offset + middle])
                if len(candidate) <= max_bytes:
                    best = (middle, candidate)
                    low = middle + 1
                else:
                    high = middle - 1
            if best is None:
                raise ValueError('one import row and its official template exceed the file size limit')
            count, data = best
        path = output / f'{prefix}_{len(files) + 1:03d}.xlsx'
        if path.exists() and path.read_bytes() != data:
            raise ValueError('export destination contains a different workbook; use a new output directory')
        if not path.exists():
            temporary = path.with_suffix('.xlsx.tmp')
            temporary.write_bytes(data)
            temporary.replace(path)
        files.append({'path': str(path.resolve()), 'rows': count, 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(), 'worksheetPart': part})
        offset += count
    return files

def _image_url(value):
    parsed = urlparse(str(value or ''))
    host = (parsed.hostname or '').lower()
    return bool(parsed.scheme == 'https' and (not parsed.username) and (not parsed.password) and parsed.path and any((host == domain or host.endswith('.' + domain) for domain in ('jd.com', '360buyimg.com'))))

def export_records(records, output, *, material_template=MATERIAL_TEMPLATE, title_template=TITLE_TEMPLATE):
    import jd_material_agent as agent
    output = Path(output)
    for template, sheet_name, width in ((material_template, '通用场景', 15), (title_template, '批量维护短标题', 2)):
        workbook = load_workbook(template, read_only=True)
        try:
            if sheet_name not in workbook.sheetnames or workbook[sheet_name].max_column != width:
                raise ValueError('official import template has an unexpected sheet or column count')
            if width == 2 and tuple(next(workbook[sheet_name].values)) != ('skuId', '短标题'):
                raise ValueError('short-title template headers must remain skuId and 短标题')
        finally:
            workbook.close()
    materials, titles, incomplete, scope = ([], [], [], [])
    seen_spus, seen_skus = (set(), set())
    for item in records:
        spu = str(item['spuId'])
        if not spu.isdecimal() or spu in seen_spus:
            raise ValueError('duplicate or invalid SPU in export scope')
        seen_spus.add(spu)
        sku_ids = list(dict.fromkeys(item['skuIds']))
        if not sku_ids or any((not str(sku).isdecimal() for sku in sku_ids)):
            raise ValueError('invalid SKU scope')
        scope.append({'spuId': spu, 'skuIds': sku_ids, 'sourceFiles': item.get('sourceFiles', [])})
        required = item.get('requiredKinds', [])
        points_required = item.get('needsSellingPoints', False)
        if required or points_required:
            points = item.get('sellingPoints', [])
            urls = item.get('urls', {})
            missing = [kind for kind in required if not _image_url(urls.get(kind))]
            if points_required and (len(points) != 3 or any((not str(point).strip() for point in points))):
                missing.append('selling_points')
            if missing:
                incomplete.append({'spuId': spu, 'skuIds': sku_ids, 'kind': 'materials', 'reason': '缺少已确认的素材或图片空间链接: ' + ','.join(missing)})
            else:
                materials.append([spu, '', *(points if points_required else ['', '', '']), '', '', '', '', '', *(urls[kind] if kind in required else '' for kind in IMAGE_COLUMNS)])
        for sku in item.get('shortTitleSkuIds', []):
            if sku not in sku_ids or sku in seen_skus:
                raise ValueError('duplicate or out-of-scope SKU short title')
            seen_skus.add(sku)
            title = item.get('shortTitles', {}).get(sku)
            if agent.valid_short_title(title):
                titles.append([sku, title])
            else:
                incomplete.append({'spuId': spu, 'skuIds': [sku], 'kind': 'shortTitle', 'reason': '短标题未完成或格式无效'})
    material_files = write_parts(materials, material_template, '通用场景', output, '商品素材导入模板-通用场景', max_rows=10000, max_bytes=MATERIAL_MAX_BYTES)
    title_files = write_parts(titles, title_template, '批量维护短标题', output, '批量维护短标题', max_bytes=TITLE_MAX_BYTES)
    return {'status': 'ready-for-manual-import' if not incomplete else 'ready-for-manual-import-with-incomplete', 'materialFiles': material_files, 'shortTitleFiles': title_files, 'materialRows': len(materials), 'shortTitleRows': len(titles), 'scope': scope, 'incomplete': incomplete, 'bindingPerformed': False, 'importSubmitted': False, 'manualImportRequired': True}

def collect_cached_records(source_dirs, erp, *, owner_erp='', agent_module=None):
    if agent_module is None:
        import jd_material_agent as agent_module
    records, sources = ({}, [])
    plans = sorted({path.resolve() for directory in source_dirs for path in Path(directory).glob('**/.state/self-operated-plan.json')})
    if not plans:
        raise ValueError('no frozen ERP plan found in the selected source directories')
    candidates = []
    for plan_path in plans:
        plan = json.loads(plan_path.read_text(encoding='utf-8'))
        folder = plan_path.parent.parent
        result = agent_module.load_self_operated_plan(folder, erp=erp, owner_erp=owner_erp, target=plan['target'], token=plan['confirmToken'])
        sources.append({'path': str(plan_path), 'sha256': hashlib.sha256(plan_path.read_bytes()).hexdigest()})
        for index, rows in enumerate(agent_module.split_rows_by_spu(result.prepared.rows, maximum_spus=50), 1):
            batch = agent_module._subset_prepared(result.prepared, rows)
            state_path = folder / f'批次{index:03d}' / '.state/task.json'
            state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
            ledger_path = state_path.with_name('upload-ledger.json')
            ledger = json.loads(ledger_path.read_text(encoding='utf-8')).get('files', {}) if ledger_path.exists() else {}
            stamp = state_path.stat().st_mtime_ns if state_path.exists() else 0
            candidates.append((stamp, str(state_path), batch, state, ledger, result.existing_materials))
    for _, source, batch, state, ledger, existing in sorted(candidates, key=lambda item: (item[0], item[1])):
        for job in batch.jobs:
            required = agent_module._model_kinds(job) + (['transparent'] if job.needs_white else [])
            if not required and (not job.needs_selling_points) and (not job.short_title_sku_ids):
                continue
            current = records.setdefault(job.spu_id, {'spuId': job.spu_id, 'skuIds': list(job.sku_ids), 'requiredKinds': [], 'needsSellingPoints': False, 'sellingPoints': [], 'urls': {}, 'shortTitleSkuIds': [], 'shortTitles': {}, 'sourceFiles': []})
            if set(current['skuIds']) != set(job.sku_ids):
                raise ValueError('SPU SKU membership changed across saved plans')
            current['sourceFiles'].append(source)
            item = state.get('jobs', {}).get(job.spu_id, {})
            known = existing.get(job.spu_id, {}).get('uploaded_urls', {})
            for kind in required:
                if kind not in current['requiredKinds']:
                    current['requiredKinds'].append(kind)
                url = item.get('uploaded_urls', {}).get(kind)
                receipt = ledger.get(agent_module.image_names(job.spu_id)[kind], {})
                if url and (receipt.get('status') == 'verified' and receipt.get('url') == url or known.get(kind) == url):
                    current['urls'][kind] = url
                elif item:
                    current['urls'].pop(kind, None)
            if job.needs_selling_points:
                current['needsSellingPoints'] = True
                if item.get('selling_points'):
                    current['sellingPoints'] = item['selling_points']
            for sku in job.short_title_sku_ids:
                if sku not in current['shortTitleSkuIds']:
                    current['shortTitleSkuIds'].append(sku)
                if item.get('short_titles', {}).get(sku):
                    current['shortTitles'][sku] = item['short_titles'][sku]
    return (list(records.values()), sources)

def export_cached_imports(source_dirs, output, erp, *, owner_erp='', agent_module=None):
    records, sources = collect_cached_records(source_dirs, erp, owner_erp=owner_erp, agent_module=agent_module)
    report = export_records(records, output)
    report.update({'erp': erp, 'sourcePlans': sources, 'newImageUploads': 0, 'modelCalls': 0})
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    (output / '导出清单.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if report['incomplete']:
        with (output / '未就绪明细.csv').open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(['SPUID', 'SKUID', '内容', '原因'])
            for item in report['incomplete']:
                writer.writerows(([item['spuId'], sku, item['kind'], item['reason']] for sku in item['skuIds']))
    (output / '上传说明.txt').write_text(f"素材表：{report['materialRows']} 个SPU；短标题表：{report['shortTitleRows']} 个SKU。\n素材表上传到 https://osw.jd.com/materialCenter/materialScene?businessModel=jxpop&platform=erp 的批量导入。\n短标题表上传到 https://osw.jd.com/site-fe/batch-task/create 的批量维护短标题。\n素材表每SPU一行，SKU列留空，作用于该SPU下全部SKU；只填写计划维护的类型，其余字段按插件留空。\n素材表每份不超过10000个SPU且不超过1MiB；短标题表每份小于5MiB，无额外10000行限制。\n这些是本地上传文件，尚未提交导入。来源缓存可能包含此前已提交的同值素材；本次未查询后台。\n缺少图片空间链接的SPU不混入可上传素材表，未就绪明细保留全部SKU。不要上传未就绪明细。\n", encoding='utf-8')
    return report

def main():
    parser = argparse.ArgumentParser()
    parser.error('Use jd_material_agent.py export-pending with the frozen exclusion plan; unguarded cache export is disabled')
if __name__ == '__main__':
    raise SystemExit(main())
