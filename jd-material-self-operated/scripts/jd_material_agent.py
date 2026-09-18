from __future__ import annotations
import protected_core as _protected_core
import argparse
import base64
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import closing, contextmanager, nullcontext
from dataclasses import asdict, dataclass, replace
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import urlparse
import httpx
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps
VARIANT_ID = 'internal-self-operated'
SKILL_VERSION = '2026.09.17.1'
MAX_SPUS_PER_BATCH = 199
MAX_SPU_CONCURRENCY = 50
JDO_IMAGE_BASE = 'https://img14.360buyimg.com/imgzone/'
SELF_OPERATED_IMAGE_SLOTS = {'white': (31, 1), 'transparent': (36, 1), 'scene1': (32, 1), 'scene2': (32, 2), 'selling': (33, 1)}

@dataclass(frozen=True)
class ProviderConfig:
    id: str
    base_url: str
    text_model: str
    image_model: str
    api_key_env: str
    text_endpoint: str
    image_endpoint: str
    text_concurrency: int
    default_image_concurrency: int
    maximum_image_concurrency: int
    requests_per_minute: int | None = None
    text_start_interval_seconds: float = 0.0
    image_start_interval_seconds: float = 0.0
PROVIDER = ProviderConfig(id='jd-internal', base_url='http://llm-gw.jd.local/v1', text_model='GPT-5.6-Sol-joybuilder', image_model='GPT-image-2-joybuilder', api_key_env='JD_LLM_API_KEY', text_endpoint='responses', image_endpoint='images/edits', text_concurrency=20, default_image_concurrency=4, maximum_image_concurrency=4, requests_per_minute=None, text_start_interval_seconds=0.75, image_start_interval_seconds=1.0)
IMAGE_SUFFIXES = {'white': '白底图.jpg', 'transparent': '透明图.png', 'scene1': '场景图1.jpg', 'scene2': '场景图2.jpg', 'selling': '卖点图.jpg'}
BANNED_SELLING_POINT_TERMS = {'第一', '顶级', '独家', '国家级', '最先进', '金牌', '质量免检', '销量', '好评', '包邮', '京东', '自营', '物流', '官方', '授权', '专供', '特供', '热卖', '疯抢', '直降', '清仓', '推荐', '爆款', '首发', '让利', '特价', '折', '赠'}

@dataclass(frozen=True)
class SourceRow:
    sku_id: str
    title: str
    spu_id: str
    score: str = ''
    raw: tuple[object, ...] = ()
    source_path: str = ''
    sheet_name: str = ''
    row_number: int = 0
    header_row: int = 1

@dataclass(frozen=True)
class SpuJob:
    spu_id: str
    representative_sku_id: str
    title: str
    sku_ids: tuple[str, ...]
    skus: tuple[tuple[str, str], ...]
    short_title_sku_ids: tuple[str, ...] = ()
    needs_selling_points: bool = True
    needs_white: bool = True
    needs_scenes: bool = True
    needs_selling_image: bool = True

    @property
    def sku_count(self) -> int:
        return len(self.sku_ids)

@dataclass(frozen=True)
class PreparedSource:
    source_type: str
    sheet_name: str
    headers: tuple[object, ...]
    rows: tuple[SourceRow, ...]
    jobs: tuple[SpuJob, ...]
    original_rows: int
    filtered_rows: int
    duplicate_rows: int

@dataclass(frozen=True)
class SelfOperatedDiscoveryResult:
    prepared: PreparedSource
    reference_urls: dict[str, str]
    existing_materials: dict[str, dict[str, Any]]
    report_rows: tuple[dict[str, Any], ...]
    failed_spus: tuple[str, ...]
    summary: dict[str, Any]
    shop_name: str
    masked_shop_id: str
    safe_to_continue: bool

@dataclass(frozen=True)
class BatchResult:
    total_spus: int
    completed_spus: int
    failed_spus: int
    output_dir: Path
    manifest_path: Path | None
    short_title_path: Path | None
    failure_path: Path | None

@dataclass(frozen=True)
class UploadResult:
    uploaded_files: int
    reused_files: int
    failed_files: int
    browser_required: bool
    url_export_path: Path
    failures: tuple[str, ...] = ()
    failed_spu_ids: tuple[str, ...] = ()

class WebcliTransportError(RuntimeError):
    pass

def cell_text(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def image_names(spu_id: str) -> dict[str, str]:
    text = cell_text(spu_id)
    if not re.fullmatch('\\d+', text):
        raise ValueError('SPUID 必须是数字文本。')
    return {kind: f'{text}{suffix}' for kind, suffix in IMAGE_SUFFIXES.items()}

def weighted_length(value: object) -> int:
    return sum((1 if ord(character) < 128 else 2 for character in cell_text(value)))

def valid_short_title(value: object) -> bool:
    text = cell_text(value)
    return bool(text) and 16 <= weighted_length(text) <= 30

def validate_selling_points(points: Iterable[object]) -> list[str]:
    normalized = [cell_text(point) for point in points]
    if len(normalized) != 3 or len(set(normalized)) != 3:
        raise ValueError('卖点必须为 3 个互不重复的词组。')
    for point in normalized:
        if not 2 <= len(point) <= 8:
            raise ValueError(f'卖点必须为 2-8 个字符：{point}')
        if re.search('[★☞【】/|｜▏♥♡]', point):
            raise ValueError(f'卖点包含无关符号：{point}')
        banned = next((term for term in BANNED_SELLING_POINT_TERMS if term in (re.sub('折叠|弯折|折弯|对折挤水', '', point) if term == '折' else point)), None)
        if banned:
            raise ValueError(f'卖点包含禁用词“{banned}”：{point}')
    return normalized

def _numeric_score(value: object) -> float | None:
    try:
        return float(cell_text(value)) if cell_text(value) else None
    except ValueError:
        return None

def collect_jdo_pages(fetch_page: Callable[[int, int], dict[str, Any]], page_size: int, id_field: str, label: str) -> list[dict[str, Any]]:
    if page_size < 1:
        raise ValueError('page_size must be positive')
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    expected_total: int | None = None
    page = 1
    while expected_total is None or len(rows) < expected_total:
        payload = fetch_page(page, page_size)
        current_page = int(payload.get('pageNum') or payload.get('currentPage') or page)
        if current_page != page:
            raise ValueError(f'{label} pagination is not contiguous: expected page {page}, got {current_page}')
        total = int(payload.get('totalItems') or 0)
        if expected_total is None:
            expected_total = total
        elif total != expected_total:
            raise ValueError(f'{label} total changed during pagination: {expected_total} -> {total}')
        items = payload.get('items') or []
        if not isinstance(items, list) or len(items) > page_size:
            raise ValueError(f'{label} returned an invalid page')
        if not items and len(rows) < expected_total:
            raise ValueError(f'{label} pagination ended before totalItems')
        for item in items:
            identifier = cell_text(item.get(id_field))
            if not identifier:
                raise ValueError(f'{label} contains an empty {id_field}')
            if identifier in seen:
                raise ValueError(f'{label} contains duplicate {id_field}: {identifier}')
            seen.add(identifier)
            rows.append(item)
        if len(rows) > expected_total:
            raise ValueError(f'{label} returned more rows than totalItems')
        page += 1
    if len(rows) != expected_total or len(seen) != expected_total:
        raise ValueError(f'{label} completeness check failed')
    return rows

class AgentBrowserSelfOperatedClient:
    page_url = 'https://osw.jd.com/materialCenter/materialScene?businessModel=jxpop&platform=erp'
    image_space_url = 'https://imgzone.shop.jd.com/?hideMenu=1&isErpSpace=1&businessModel=jxpop'

    def __init__(self, *, cdp: str | None=None, profile: str | None=None, session: str='jd-self-operated-readonly', timeout: float=180.0):
        self.cdp = cell_text(cdp)
        self.profile = cell_text(profile)
        self.session = session
        self.image_session = f'{session}-image-space'
        self.timeout = timeout
        self.initialized = False
        self.image_space_initialized = False
        self.expected_shop_name = ''
        self.bridge_path = Path(__file__).with_name('self_operated_browser_bridge.js')
        self.upload_bridge_path = Path(__file__).with_name('self_operated_image_upload_bridge.js')
        self.binding_bridge_path = Path(__file__).with_name('self_operated_material_binding_bridge.js')
        executable = shutil.which('agent-browser')
        if executable:
            self.command = [executable]
        else:
            npx = shutil.which('npx') or shutil.which('npx.cmd')
            if not npx:
                raise RuntimeError('agent-browser is unavailable')
            self.command = [npx, '--yes', 'agent-browser']

    def _base(self, *, image_space: bool=False) -> list[str]:
        session = self.image_session if image_space else self.session
        command = [*self.command, '--session', session, '--pin-tab', '--json']
        if self.cdp:
            command.extend(('--cdp', self.cdp))
        elif self.profile:
            command.extend(('--profile', self.profile))
        else:
            command.append('--auto-connect')
        return command

    def _run(self, arguments: list[str], input_text: str | None=None, *, image_space: bool=False) -> Any:
        completed = subprocess.run([*self._base(image_space=image_space), *arguments], capture_output=True, timeout=self.timeout, check=False, input=input_text.encode('utf-8') if input_text is not None else None)
        output = None
        for encoding in ('utf-8', 'gb18030'):
            try:
                output = completed.stdout.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            payload = json.loads((output or '').strip())
        except json.JSONDecodeError as error:
            raise RuntimeError('authenticated browser bridge did not return JSON') from error
        if completed.returncode or payload.get('success') is False:
            raise RuntimeError('authenticated browser bridge failed; check login or security verification')
        data = payload.get('data')
        if isinstance(data, dict) and 'result' in data:
            return data['result']
        return data

    def _eval(self, action: dict[str, Any]) -> Any:
        if not self.initialized:
            current = self._run(['get', 'url'])
            current_url = cell_text((current or {}).get('url') if isinstance(current, dict) else current)
            if not current_url.startswith('https://osw.jd.com/'):
                self._run(['open', self.page_url])
            self.initialized = True
            self.image_space_initialized = False
        source = self.bridge_path.read_text(encoding='utf-8')
        source = source.replace('__JD_SELF_OPERATED_ACTION__', json.dumps(action, ensure_ascii=False))
        return self._run(['eval', '--stdin'], input_text=source)

    def _eval_image_space(self, action: dict[str, Any]) -> Any:
        if not self.image_space_initialized:
            current = self._run(['get', 'url'], image_space=True)
            current_url = cell_text((current or {}).get('url') if isinstance(current, dict) else current)
            if not current_url.startswith('https://imgzone.shop.jd.com/'):
                self._run(['open', self.image_space_url], image_space=True)
            self.image_space_initialized = True
        source = self.upload_bridge_path.read_text(encoding='utf-8')
        source = source.replace('__JD_SELF_OPERATED_UPLOAD_ACTION__', json.dumps(action, ensure_ascii=False))
        return self._run(['eval', '--stdin'], input_text=source, image_space=True)

    def select_shop(self, expected_shop_name: str | None=None) -> dict[str, Any]:
        data = self._eval({'kind': 'initialize', 'expectedShopName': expected_shop_name or '', 'expectedErp': getattr(self, 'expected_erp', '')})
        if not isinstance(data, dict):
            raise ValueError('shop enumeration returned an invalid response')
        self.expected_shop_name = cell_text(data.get('shopName'))
        return data

    def spu_page(self, page: int, page_size: int) -> dict[str, Any]:
        data = self._eval({'kind': 'spuPage', 'page': int(page), 'pageSize': int(page_size)})
        if not isinstance(data, dict):
            raise ValueError('SPU listing returned an invalid response')
        return data

    def target_spu_page(self, spu_id: str, page: int, page_size: int) -> dict[str, Any]:
        data = self._eval({'kind': 'spuPage', 'page': int(page), 'pageSize': int(page_size), 'spuIds': [spu_id]})
        if not isinstance(data, dict):
            raise ValueError('target SPU enumeration returned an invalid response')
        return data

    def inspect_spus(self, spu_ids: Iterable[str]) -> list[dict[str, Any]]:
        data = self._eval({'kind': 'inspectSpus', 'spuIds': [str(value) for value in spu_ids]})
        if not isinstance(data, list):
            raise ValueError('SPU inspection returned an invalid response')
        return data

    def inspect_material_details(self, spu_ids: Iterable[str]) -> list[dict[str, Any]]:
        rows = []
        for group in _chunks([str(value) for value in spu_ids], 5):
            data = self._eval({'kind': 'inspectSpuDetails', 'spuIds': group})
            if not isinstance(data, list) or len(data) != len(group) or {cell_text(row.get('spuId')) for row in data} != set(group):
                raise ValueError('material detail inspection returned an incomplete SPU scope')
            rows.extend(data)
        return rows

    def image_page(self, category_id: int, page: int, page_size: int, query: str='') -> dict[str, Any]:
        data = self._eval_image_space({'kind': 'imagePage', 'categoryId': str(category_id), 'page': int(page), 'pageSize': int(page_size), 'query': cell_text(query)})
        if not isinstance(data, dict):
            raise ValueError('image-space listing returned an invalid response')
        return data

    def upload_image(self, path: str | Path, category_id: int) -> dict[str, Any]:
        image_path = Path(path)
        data = self._eval_image_space({'kind': 'uploadImage', 'categoryId': str(category_id), 'fileName': image_path.name, 'fileData': base64.b64encode(image_path.read_bytes()).decode('ascii')})
        if not isinstance(data, dict):
            raise ValueError('image-space upload returned an invalid response')
        return data

class WebcliSelfOperatedClient(AgentBrowserSelfOperatedClient):

    def __init__(self, *, profile: str, session: str='jd-self-operated-webcli', timeout: float=180.0):
        self.command = list(_webcli_command())
        self.profile = cell_text(profile)
        self.session = session
        self.image_session = f'{session}-image-space'
        self.timeout = timeout
        self.initialized = False
        self.image_space_initialized = False
        self.expected_shop_name = ''
        self.bridge_path = Path(__file__).with_name('self_operated_browser_bridge.js')
        self.binding_bridge_path = Path(__file__).with_name('self_operated_material_binding_bridge.js')
        self.upload_bridge_path = Path(__file__).with_name('self_operated_image_upload_bridge.js')
        self.product_bridge_path = Path(__file__).with_name('self_operated_product_bridge.js')
        self.product_session = f'{session}-products'
        self.product_initialized = False
        self.expected_erp = ''
        self.owner_erp = ''
        self.merge_image_types = True
        self.binding_concurrency = 5
        self.binding_batch_size = 5
        self.upload_concurrency = 5
        self.upload_batch_size = 5
        self.short_title_target_rows = 50
        self.cache_bridges = True
        self.pipeline_enabled = True

    def _run_webcli(self, arguments: list[str], *, image_space: bool=False) -> Any:
        session = self.image_session if image_space else self.session
        completed = subprocess.run([*self.command, '--profile', self.profile, 'browser', session, '--window', 'background', *arguments], capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=self.timeout, check=False)
        try:
            payload = json.loads(completed.stdout.strip())
        except json.JSONDecodeError as error:
            if any((re.match('^\\s*(?:✖\\s*)?Error: authenticated ERP', line) for line in (completed.stderr or '').splitlines())):
                raise RuntimeError('authenticated ERP verification failed; browser task stopped') from error
            raise WebcliTransportError(f'WebCLI browser bridge did not return JSON (exit {completed.returncode})') from error
        native_failure = arguments[:1] != ['eval'] and isinstance(payload, dict) and (payload.get('ok') is False)
        if completed.returncode or native_failure:
            detail = cell_text(payload.get('error') if isinstance(payload, dict) else '')
            raise RuntimeError(detail or 'WebCLI browser bridge failed; check login or security verification')
        return payload

    def _ensure_browser_page(self, url, *, image_space=False):
        source = f'(()=>{{const expected=new URL({json.dumps(url)});const current=new URL(location.href);return current.origin===expected.origin&&current.pathname===expected.pathname&&Array.from(expected.searchParams).every(([name,value])=>current.searchParams.get(name)===value)}})()'
        try:
            ready = self._run_webcli(['eval', source], image_space=image_space) is True
        except _protected_core.ProtectedCoreError:
            raise
        except RuntimeError:
            ready = False
        if not ready:
            self._run_webcli(['open', url], image_space=image_space)

    def _eval(self, action: dict[str, Any]) -> Any:
        if not self.initialized or action.get('kind') == 'initialize':
            self._ensure_browser_page(self.page_url)
            self.initialized = True
            self.image_space_initialized = False
        source = self.bridge_path.read_text(encoding='utf-8')
        for attempt in range(2):
            result = self._eval_template(source, '__JD_SELF_OPERATED_ACTION__', action, readonly=True, recovery_url=self.page_url)
            context_missing = action.get('kind') in {'inspectSpus', 'inspectSpuDetails'} and isinstance(result, list) and bool(result) and all((isinstance(row, dict) and (not row.get('skus')) and (not row.get('textChildren')) and (row.get('errors') == ['self-operated shop is not initialized']) for row in result))
            if not context_missing:
                return result
            if attempt:
                raise RuntimeError('self-operated material context is unavailable after reinitialization')
            self.select_shop(self.expected_shop_name or None)

    def _eval_image_space(self, action: dict[str, Any]) -> Any:
        if action.get('kind') == 'uploadImage' and getattr(self, 'expected_erp', ''):
            self.select_shop()
        if not self.image_space_initialized:
            self._ensure_browser_page(self.image_space_url, image_space=True)
            self.image_space_initialized = True
        return self._eval_template(self.upload_bridge_path.read_text(encoding='utf-8'), '__JD_SELF_OPERATED_UPLOAD_ACTION__', {**action, 'expectedErp': getattr(self, 'expected_erp', '')}, image_space=True, readonly=action.get('kind') == 'imagePage', recovery_url=self.image_space_url)

    def _eval_readonly_webcli_source(self, source: str, *, image_space: bool=False, recovery_url: str='') -> Any:
        for attempt in range(3):
            try:
                return self._eval_webcli_source(source, image_space=image_space)
            except _protected_core.ProtectedCoreError:
                raise
            except WebcliTransportError:
                if attempt == 2:
                    raise
                if recovery_url:
                    self._ensure_browser_page(recovery_url, image_space=image_space)
                time.sleep(attempt + 1)
            except RuntimeError as error:
                if attempt or 'command_result_unknown' not in str(error):
                    raise
                time.sleep(1)

    def _eval_template(self, template, marker, action, *, image_space=False, readonly=False, recovery_url=''):
        if not getattr(self, 'cache_bridges', False):
            source = template.replace(marker, json.dumps(action, ensure_ascii=False))
            if readonly:
                return self._eval_readonly_webcli_source(source, image_space=image_space, recovery_url=recovery_url)
            return self._eval_webcli_source(source, image_space=image_space)
        for attempt in range(3 if readonly else 1):
            try:
                return self._eval_cached_template(template, marker, action, image_space=image_space)
            except _protected_core.ProtectedCoreError:
                raise
            except WebcliTransportError:
                if not readonly or attempt == 2:
                    raise
                if recovery_url:
                    self._ensure_browser_page(recovery_url, image_space=image_space)
                time.sleep(attempt + 1)
            except RuntimeError as error:
                if not readonly or attempt or 'command_result_unknown' not in str(error):
                    raise
                time.sleep(1)

    def _eval_cached_template(self, template, marker, action, *, image_space=False):
        key = '__jdMaterialBridge_' + hashlib.sha256(template.encode('utf-8')).hexdigest()
        encoded = json.dumps(key)
        payload = json.dumps(action, ensure_ascii=False)
        invocation = f"(async()=>{{if(typeof window[{encoded}]!=='function')return {{bridgeCacheMiss:true}};return {{bridgeResult:await window[{encoded}]({payload})}}}})()"
        reply = self._eval_webcli_source(invocation, image_space=image_space)
        if isinstance(reply, dict) and reply.get('bridgeCacheMiss') is True:
            body = template.replace(marker, 'bridgeAction').rstrip().removesuffix(';')
            self._eval_webcli_source(f'window[{encoded}]=(bridgeAction)=>{body};true', image_space=image_space)
            reply = self._eval_webcli_source(invocation, image_space=image_space)
        if not isinstance(reply, dict) or 'bridgeResult' not in reply:
            raise RuntimeError('cached business bridge returned no result; do not replay a write')
        return reply['bridgeResult']

    def _eval_webcli_source(self, source: str, *, image_space: bool=False) -> Any:
        if len(source.encode('utf-8')) < 12000:
            return self._run_webcli(['eval', source], image_space=image_space)
        key = '__jdMaterialScript_' + uuid.uuid4().hex
        encoded = json.dumps(key)
        try:
            self._run_webcli(['eval', f'window[{encoded}]=[]; true'], image_space=image_space)
            for offset in range(0, len(source), 3000):
                chunk = json.dumps(source[offset:offset + 3000], ensure_ascii=True)
                self._run_webcli(['eval', f'window[{encoded}].push({chunk}); true'], image_space=image_space)
            return self._run_webcli(['eval', f"(0,eval)(window[{encoded}].join(''))"], image_space=image_space)
        finally:
            try:
                self._run_webcli(['eval', f'delete window[{encoded}]'], image_space=image_space)
            except _protected_core.ProtectedCoreError:
                raise
            except Exception:
                pass

    def _clear_upload_input(self):
        input_id = getattr(self, '_upload_input_id', None)
        self._upload_input_id = None
        self._upload_staged_paths = set()
        if input_id:
            try:
                self._run_webcli(['eval', f'document.getElementById({json.dumps(input_id)})?.remove();true'], image_space=True)
            except _protected_core.ProtectedCoreError:
                raise
            except Exception:
                pass

    @contextmanager
    def stage_upload_files(self, paths):
        if getattr(self, '_upload_candidates', None) is not None:
            raise ValueError('nested upload staging is not supported')
        self._upload_candidates = tuple((Path(path).resolve() for path in paths))
        try:
            yield
        finally:
            self._clear_upload_input()
            self._upload_candidates = None

    def upload_image(self, path: str | Path, category_id: int) -> dict[str, Any]:
        path = Path(path).resolve()
        candidates = getattr(self, '_upload_candidates', None)
        try:
            if str(path) not in getattr(self, '_upload_staged_paths', set()):
                self._clear_upload_input()
                if not self.image_space_initialized:
                    self._ensure_browser_page(self.image_space_url, image_space=True)
                    self.image_space_initialized = True
                selected = [path]
                prefix = re.match('^\\d+', path.name)
                if candidates is not None and prefix:
                    for candidate in candidates:
                        sibling = re.match('^\\d+', candidate.name)
                        if sibling and sibling.group(0) == prefix.group(0) and candidate.is_file() and (candidate.name not in {item.name for item in selected}):
                            selected.append(candidate)
                        if len(selected) == 5:
                            break
                self._upload_input_id = 'jd-material-upload-' + uuid.uuid4().hex
                encoded = json.dumps(self._upload_input_id)
                self._run_webcli(['eval', f"(()=>{{const input=document.createElement('input');input.type='file';input.multiple=true;input.id={encoded};input.hidden=true;document.body.appendChild(input);return true}})()"], image_space=True)
                self._run_webcli(['upload', f'#{self._upload_input_id}', *[str(item) for item in selected]], image_space=True)
                self._upload_staged_paths = {str(item) for item in selected}
            return self._eval_image_space({'kind': 'uploadImage', 'categoryId': str(category_id), 'fileName': path.name, 'fileInputId': self._upload_input_id, 'expectedSha256': hashlib.sha256(path.read_bytes()).hexdigest()})
        finally:
            if candidates is None:
                self._clear_upload_input()

    def upload_images(self, paths, category_id, *, expected_hashes=None):
        paths = [Path(path).resolve() for path in paths]
        prefixes = [re.match('^\\d+', path.name) for path in paths]
        if not self.expected_erp or not paths or any((prefix is None for prefix in prefixes)) or (len({path.name for path in paths}) != len(paths)):
            raise ValueError('batch upload requires an ERP and distinct authorized product files')
        if expected_hashes is not None and (set(expected_hashes) != {path.name for path in paths} or any((not isinstance(value, str) or not re.fullmatch('[a-f0-9]{64}', value) for value in expected_hashes.values()))):
            raise ValueError('batch upload hashes must cover the exact approved file set')
        files = [{'fileName': path.name, 'expectedSha256': expected_hashes[path.name] if expected_hashes is not None else _file_sha256(path)} for path in paths]
        self._clear_upload_input()
        try:
            if not self.image_space_initialized:
                self._ensure_browser_page(self.image_space_url, image_space=True)
                self.image_space_initialized = True
            self._upload_input_id = 'jd-material-upload-' + uuid.uuid4().hex
            encoded = json.dumps(self._upload_input_id)
            self._run_webcli(['eval', f"(()=>{{const input=document.createElement('input');input.type='file';input.multiple=true;input.id={encoded};input.hidden=true;document.body.appendChild(input);return true}})()"], image_space=True)
            self._run_webcli(['upload', f'#{self._upload_input_id}', *[str(path) for path in paths]], image_space=True)
            source = Path(__file__).with_name('self_operated_image_batch_upload_bridge.js').read_text(encoding='utf-8')
            action = {'expectedErp': self.expected_erp, 'categoryId': str(category_id), 'fileInputId': self._upload_input_id, 'files': files, 'concurrency': getattr(self, 'upload_concurrency', 1)}
            single = self.upload_bridge_path.read_text(encoding='utf-8').replace('__JD_SELF_OPERATED_UPLOAD_ACTION__', 'singleAction')
            source = source.replace('__JD_SELF_OPERATED_UPLOAD_SINGLE_SOURCE__', single)
            return self._eval_template(source, '__JD_SELF_OPERATED_UPLOAD_BATCH_ACTION__', action, image_space=True)
        finally:
            self._clear_upload_input()

    def _eval_products(self, action: dict[str, Any]) -> Any:
        client = WebcliSelfOperatedClient(profile=self.profile, session=self.product_session, timeout=self.timeout)
        client.cache_bridges = getattr(self, 'cache_bridges', False)
        page_url = 'https://wares-jdm.jd.com/ware/wareList?businessModel=2&loginType=2'
        if not self.product_initialized:
            client._ensure_browser_page(page_url)
            self.product_initialized = True
        return client._eval_template(self.product_bridge_path.read_text(encoding='utf-8'), '__JD_SELF_OPERATED_PRODUCT_ACTION__', {**action, 'expectedErp': self.expected_erp, 'ownerErp': self.owner_erp}, readonly=action.get('kind') in {'inventoryPage', 'shortTitleRows'}, recovery_url=page_url)

    def inventory_page(self, page: int, page_size: int) -> dict[str, Any]:
        return self._eval_products({'kind': 'inventoryPage', 'page': page, 'pageSize': page_size})

    def target_inventory_page(self, spu_id: str) -> dict[str, Any]:
        return self._eval_products({'kind': 'inventoryPage', 'page': 1, 'pageSize': 1, 'spuIds': [spu_id]})

    def scoped_inventory_page(self, spu_ids: list[str]) -> dict[str, Any]:
        return self._eval_products({'kind': 'inventoryPage', 'page': 1, 'pageSize': 100, 'spuIds': spu_ids})

    def short_title_rows(self, spu_ids: Iterable[str]) -> list[dict[str, Any]]:
        rows = []
        for group in _chunks(list(spu_ids), 5):
            values = self._eval_products({'kind': 'shortTitleRows', 'spuIds': group})
            if not isinstance(values, list):
                raise ValueError('SKU readback returned an invalid batch')
            rows.extend(values)
        return rows

    def verify_product_scope(self, jobs):
        if not jobs:
            return []
        expected = {job.spu_id: set(job.sku_ids) for job in jobs}
        payload = self._eval_products({'kind': 'inventoryPage', 'page': 1, 'pageSize': 100, 'spuIds': list(expected)})
        if payload.get('page') != 1 or payload.get('total') != len(expected) or len(payload.get('rows', [])) != len(expected) or ({cell_text(row.get('spuId')) for row in payload.get('rows', [])} != set(expected)):
            raise ValueError('authorized product scope changed since confirmation')
        observed = {spu_id: set() for spu_id in expected}
        rows = self.short_title_rows(list(expected))
        for row in rows:
            spu_id, sku_id = (cell_text(row['productId']), cell_text(row['skuId']))
            if spu_id not in observed or sku_id in observed[spu_id]:
                raise ValueError('unexpected or duplicate SKU in current scope')
            observed[spu_id].add(sku_id)
        if observed != expected:
            raise ValueError('SKU scope changed since confirmation; create a new plan')
        return rows

def _webcli_command() -> tuple[str, ...]:
    executable = shutil.which('webcli')
    if not executable:
        raise RuntimeError('webcli is unavailable')
    executable_path = Path(executable)
    if executable_path.suffix.lower() in {'.cmd', '.bat', '.ps1'}:
        node = shutil.which('node')
        entry = executable_path.parent / 'node_modules' / '@jd' / 'webcli' / 'dist' / 'src' / 'main.js'
        if node and entry.is_file():
            roots = (Path(__file__).resolve().parents[2], Path.home() / '.agents' / 'skills', Path.home() / '.codex' / 'skills')
            overlay = next((root / 'webcli-browser-runtime' / 'scripts' / 'webcli-preload.mjs' for root in roots if (root / 'webcli-browser-runtime' / 'scripts' / 'webcli-preload.mjs').is_file()), None)
            if overlay is None:
                raise RuntimeError('Codex must install the companion webcli-browser-runtime Skill before browser maintenance')
            return (node, '--import', overlay.resolve().as_uri(), str(entry))
    return (executable,)

def _connected_webcli_profiles(timeout: float=30.0) -> tuple[str, ...]:
    command = _webcli_command()
    completed = subprocess.run([*command, '--json', 'doctor'], capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=timeout, check=False)
    try:
        payload = json.loads(completed.stdout.strip())
    except json.JSONDecodeError:
        return ()
    if payload.get('code') == 'WEBCLI_RUNTIME_MERGE_REQUIRED':
        raise RuntimeError(cell_text(payload.get('error')) or 'WebCLI runtime requires upstream merge')
    if completed.returncode:
        return ()
    return tuple((cell_text(item.get('contextId')) for item in payload.get('profiles', []) if item.get('extensionConnected') and cell_text(item.get('contextId'))))

def create_self_operated_browser_client(args) -> AgentBrowserSelfOperatedClient:
    browser_cdp = getattr(args, 'browser_cdp', None)
    browser_profile = getattr(args, 'browser_profile', None)
    webcli_profile = getattr(args, 'webcli_profile', None)
    if sum((bool(value) for value in (browser_cdp, browser_profile, webcli_profile))) > 1:
        raise ValueError('choose only one of --browser-cdp, --browser-profile, or --webcli-profile')
    timeout = float(getattr(args, 'timeout', 180.0))
    if webcli_profile:
        return WebcliSelfOperatedClient(profile=webcli_profile, timeout=timeout)
    if not browser_cdp and (not browser_profile):
        profiles = _connected_webcli_profiles(timeout=min(timeout, 30.0))
        if len(profiles) == 1:
            return WebcliSelfOperatedClient(profile=profiles[0], timeout=timeout)
        if len(profiles) > 1:
            erp = cell_text(getattr(args, 'erp', ''))
            if erp:
                for profile in profiles:
                    candidate = WebcliSelfOperatedClient(profile=profile, timeout=timeout)
                    candidate.expected_erp = erp
                    try:
                        candidate.select_shop()
                        return candidate
                    except _protected_core.ProtectedCoreError:
                        raise
                    except (ValueError, RuntimeError):
                        continue
                raise ValueError('请先用提供的 ERP 登录自营商品后台；未找到匹配的登录账号。')
            raise ValueError('multiple WebCLI browser profiles are connected; pass --webcli-profile')
        if not profiles and getattr(args, 'erp', None):
            raise ValueError('请先登录自营商品后台并允许 Browser Bridge 连接；Codex 负责准备工具，无需配置 CDP 或 Profile。')
    return AgentBrowserSelfOperatedClient(cdp=browser_cdp, profile=browser_profile, timeout=timeout)

def _chunks(values: list[str], size: int) -> Iterable[list[str]]:
    for start in range(0, len(values), size):
        yield values[start:start + size]

def _jd_image_url(value: object) -> str:
    url = cell_text(value)
    if not url:
        return ''
    if url.startswith('jfs/'):
        return JDO_IMAGE_BASE + url
    if url.startswith('//'):
        url = 'https:' + url
    elif not re.match('^https?://', url) and '.360buyimg.com/' in url:
        url = 'https://' + url.lstrip('/')
    parsed = urlparse(url)
    host = (parsed.hostname or '').lower()
    if parsed.scheme != 'https' or not (host.endswith('.360buyimg.com') or host.endswith('.jd.com')):
        raise ValueError('JDO returned an untrusted product image URL')
    return url

def _jfs_path(value: object) -> str:
    url = cell_text(value)
    if url.startswith('jfs/'):
        return url
    parsed = urlparse(url)
    marker = '/jfs/'
    if parsed.scheme == 'https' and marker in parsed.path:
        return parsed.path[parsed.path.index(marker) + 1:]
    raise ValueError('material image URL does not contain a relative jfs path')

def collect_self_operated_inventory(client, *, page_size: int=100) -> list[dict[str, Any]]:
    if type(page_size) is not int or page_size < 1:
        raise ValueError('authorized inventory page size must be a positive integer')

    def fetch(page: int, size: int) -> dict[str, Any]:
        payload = client.inventory_page(page, size)
        if 'total' not in payload or 'page' not in payload or (not isinstance(payload.get('rows'), list)):
            raise ValueError('authorized inventory returned an incomplete pagination contract')
        return {'pageNum': payload['page'], 'totalItems': payload['total'], 'items': payload['rows']}
    return collect_jdo_pages(fetch, page_size, 'spuId', 'self-operated authorized inventory')

def write_self_operated_plan(result, output: Path, *, erp: str, owner_erp: str, target: str) -> str:
    payload = {'formatVersion': 1, 'erp': erp, 'ownerErp': owner_erp, 'target': target, 'output': str(output.resolve()), 'result': asdict(result), 'policy': 'available-permissions-gap-triggered-spu-type-overwrite-v1'}
    token = 'SELF-' + _auto_maintain_plan_hash(payload)
    payload['confirmToken'] = token
    _save_state(output / '.state' / 'self-operated-plan.json', payload)
    return token

def load_self_operated_plan(output: Path, *, erp: str, owner_erp: str, target: str, token: str):
    payload = _load_state(output / '.state' / 'self-operated-plan.json')
    actual = payload.pop('confirmToken', '')
    if payload.get('formatVersion') != 1 or payload.get('erp') != erp or payload.get('ownerErp') != owner_erp or (payload.get('target') != target) or (payload.get('output') != str(output.resolve())):
        raise ValueError('self-operated plan identity/scope/output mismatch; create a new plan')
    if actual != token or actual != 'SELF-' + _auto_maintain_plan_hash(payload):
        raise ValueError('self-operated plan confirmation token is invalid')
    data = payload['result']
    source = dict(data['prepared'])
    source['headers'] = tuple(source['headers'])
    source['rows'] = tuple((SourceRow(**{**row, 'raw': tuple(row['raw'])}) for row in source['rows']))
    source['jobs'] = tuple((SpuJob(**{**job, 'sku_ids': tuple(job['sku_ids']), 'skus': tuple((tuple(sku) for sku in job['skus'])), 'short_title_sku_ids': tuple(job['short_title_sku_ids'])}) for job in source['jobs']))
    return SelfOperatedDiscoveryResult(**{**data, 'prepared': PreparedSource(**source), 'report_rows': tuple(data['report_rows']), 'failed_spus': tuple(data['failed_spus'])})

def create_self_operated_erp_client(args):
    client = create_self_operated_browser_client(args)
    if not isinstance(client, WebcliSelfOperatedClient):
        raise ValueError('ERP automatic workflow requires the authorized WebCLI Browser Bridge')
    client.expected_erp = cell_text(getattr(args, 'erp', ''))
    client.owner_erp = cell_text(getattr(args, 'owner_erp', ''))
    if not client.expected_erp:
        raise ValueError('provide the ERP identity; do not enter a shop name')
    scope = str(Path(args.output_dir).resolve()) if getattr(args, 'output_dir', None) else f'{client.expected_erp}|{client.owner_erp}|{self_operated_target(args)}'
    session = 'jd-material-erp-' + hashlib.sha256(scope.encode()).hexdigest()[:16]
    client.session = session
    client.image_session = f'{session}-image-space'
    client.product_session = f'{session}-products'
    client.initialized = client.image_space_initialized = client.product_initialized = False
    return client

class SelfOperatedCheckpointClient:

    def __init__(self, client, output: Path):
        self.client = client
        self.root = output / '.state' / 'discovery'
        identity = {'erp': client.expected_erp, 'ownerErp': client.owner_erp, 'source': 'products-v1'}
        path = self.root / 'identity.json'
        if path.exists() and json.loads(path.read_text(encoding='utf-8')) != identity:
            raise ValueError('discovery checkpoint identity/scope mismatch')
        _save_state(path, identity)

    def __getattr__(self, name):
        return getattr(self.client, name)

    def inventory_page(self, page, page_size):
        path = self.root / f'page-{page_size}-{page}.json'
        if path.exists():
            return json.loads(path.read_text(encoding='utf-8'))
        payload = self.client.inventory_page(page, page_size)
        _save_state(path, payload)
        return payload

    def short_title_rows(self, spu_ids):
        identifiers = list(spu_ids)
        if any((not re.fullmatch('\\d+', identifier) for identifier in identifiers)):
            raise ValueError('invalid checkpoint SPU')
        found = {}
        pending = []
        for spu_id in identifiers:
            path = self.root / f'sku-confirmation-{spu_id}.json'
            if path.is_file():
                saved = json.loads(path.read_text(encoding='utf-8'))
                if saved.get('spuId') != spu_id or not isinstance(saved.get('rows'), list):
                    raise ValueError('SKU confirmation checkpoint scope mismatch')
                found[spu_id] = saved['rows']
            else:
                pending.append(spu_id)
        if pending:
            observed = self.client.short_title_rows(pending)
            if not isinstance(observed, list) or any((not isinstance(row, dict) or cell_text(row.get('productId')) not in pending or (not re.fullmatch('\\d+', cell_text(row.get('skuId')))) for row in observed)):
                raise ValueError('SKU confirmation returned an invalid product scope')
            pairs = [(cell_text(row['productId']), cell_text(row['skuId'])) for row in observed]
            if len(set(pairs)) != len(pairs):
                raise ValueError('SKU confirmation returned duplicate scope rows')
            for spu_id in pending:
                found[spu_id] = [row for row in observed if cell_text(row['productId']) == spu_id]
                _save_state(self.root / f'sku-confirmation-{spu_id}.json', {'spuId': spu_id, 'rows': found[spu_id]})
        return [row for spu_id in identifiers for row in found[spu_id]]

    def inspect_spus(self, spu_ids):
        return self._cached_material_inspections(spu_ids, detailed=False)

    def inspect_material_details(self, spu_ids):
        return self._cached_material_inspections(spu_ids, detailed=True)

    def _cached_material_inspections(self, spu_ids, *, detailed):
        prefix = 'detail-spu' if detailed else 'spu'
        found = {}
        pending = []
        for spu_id in spu_ids:
            if not re.fullmatch('\\d+', spu_id):
                raise ValueError('invalid checkpoint SPU')
            path = self.root / f'{prefix}-{spu_id}.json'
            if path.exists():
                found[spu_id] = json.loads(path.read_text(encoding='utf-8'))
            else:
                pending.append(spu_id)
        if pending:
            inspect = self.client.inspect_material_details if detailed and callable(getattr(type(self.client), 'inspect_material_details', None)) else self.client.inspect_spus
            results = inspect(pending)
            if not isinstance(results, list):
                raise ValueError('invalid inspection response')
            seen = set()
            for result in results:
                spu_id = cell_text(result.get('spuId'))
                if spu_id not in pending or spu_id in seen:
                    raise ValueError('unexpected or duplicate inspection SPU')
                seen.add(spu_id)
                found[spu_id] = result
                if not result.get('errors'):
                    _save_state(self.root / f'{prefix}-{spu_id}.json', result)
        return [found.get(spu_id, {'spuId': spu_id, 'errors': ['lookup omitted SPU']}) for spu_id in spu_ids]

def self_operated_target(args):
    target = cell_text(getattr(args, 'target_spu_id', ''))
    selected = [cell_text(value) for value in getattr(args, 'target_spu_ids', None) or []]
    if target and selected:
        raise ValueError('choose one explicit target scope')
    if any((not re.fullmatch('\\d+', value) for value in selected)):
        raise ValueError('explicit scope requires numeric SPUs')
    if len(selected) != len(set(selected)):
        raise ValueError('duplicate SPU in explicit scope')
    if target and (not re.fullmatch('\\d+', target)):
        raise ValueError('target SPU must be numeric')
    return ','.join(sorted(selected)) if selected else target

def self_operated_output(args):
    target = self_operated_target(args)
    if not getattr(args, 'output_dir', None):
        erp = cell_text(getattr(args, 'erp', ''))
        if not re.fullmatch('[\\w.-]+', erp) or erp in {'.', '..'}:
            raise ValueError('provide a valid ERP identity')
        owner = cell_text(getattr(args, 'owner_erp', ''))
        scope = hashlib.sha256(f'{owner}|{target}'.encode()).hexdigest()[:12]
        args.output_dir = Path.home() / 'Documents' / 'JDMaterial' / 'self-operated' / erp / scope
    return Path(args.output_dir)

def run_plan_self_operated(args):
    output = self_operated_output(args)
    client = create_self_operated_erp_client(args)
    client.select_shop()
    target = self_operated_target(args)
    path = output / '.state' / 'self-operated-plan.json'
    if path.is_file():
        token = _load_state(path).get('confirmToken', '')
        result = load_self_operated_plan(output, erp=client.expected_erp, owner_erp=client.owner_erp, target=target, token=token)
        reused = True
    else:
        selected = getattr(args, 'target_spu_ids', None)
        result = discover_authorized_self_operated_source(SelfOperatedCheckpointClient(client, output), target_spu_id='' if selected else target, target_spu_ids=selected)
        token = write_self_operated_plan(result, output, erp=client.expected_erp, owner_erp=client.owner_erp, target=target)
        write_self_operated_discovery(result, output)
        reused = False
    return {'status': 'planned', 'erp': client.expected_erp, 'scope': 'explicit-proxy-filter' if client.owner_erp else 'available-permissions', 'outputDir': str(output.resolve()), 'planReused': reused, 'confirmToken': token, 'summary': result.summary, 'spuCount': len(result.prepared.jobs), 'skuCount': len(result.prepared.rows)}

def _self_image_status(sku: dict[str, Any], material_type: int, order: int) -> tuple[str, str]:
    supported = sku.get('supportedTypes')
    if not isinstance(supported, list):
        return ('unknown', '')
    if material_type not in {int(value) for value in supported}:
        return ('unsupported', '')
    matches = [item for item in sku.get('materials') or [] if int(item.get('type') or 0) == material_type and int(item.get('order') or 0) == order]
    if not matches:
        return ('missing', '')
    if any((int(item.get('status') or 0) == 3 for item in matches)):
        return ('pending', '')
    approved = [item for item in matches if int(item.get('status') or 0) == 4]
    for item in approved:
        try:
            url = _jd_image_url(item.get('url'))
        except ValueError:
            continue
        if url:
            return ('approved', url)
    if any((int(item.get('status') or 0) == 5 for item in matches)):
        return ('rejected', '')
    return ('invalid', '')

def _self_score_status(score: object, realtime_gap: bool=False) -> str:
    numeric = _numeric_score(score)
    if numeric is None:
        return 'unscored'
    if 0 <= numeric <= 89:
        return 'low'
    if 90 <= numeric <= 100:
        return 'qualified-mismatch' if realtime_gap else 'qualified'
    return 'invalid'

def _self_short_title(sku: dict[str, Any], child: dict[str, Any]) -> tuple[str, str]:
    known_values = []
    for source in (sku, child):
        if source.get('shortTitleKnown') is True:
            known_values.append(cell_text(source.get('shortTitle')))
    if not known_values:
        return ('unknown', '')
    nonempty = {value for value in known_values if value}
    if len(nonempty) > 1 or (nonempty and any((not value for value in known_values))):
        return ('conflict', '')
    value = next(iter(nonempty), '')
    return ('complete' if value else 'missing', value)

def discover_authorized_self_operated_source(client, *, target_spu_id: str='', target_spu_ids: Iterable[str] | None=None, max_spus: int | None=None) -> SelfOperatedDiscoveryResult:
    from types import SimpleNamespace
    selected = list(target_spu_ids or [])
    canonical = self_operated_target(SimpleNamespace(target_spu_id=target_spu_id, target_spu_ids=selected))
    selected = canonical.split(',') if selected else []
    target = cell_text(target_spu_id)
    if max_spus is not None and (max_spus < 1 or target or selected):
        raise ValueError('invalid diagnostic scope')
    shop = client.select_shop()
    if selected:
        payload = client.scoped_inventory_page(selected)
        rows = payload.get('rows', [])
        if payload.get('page') != 1 or payload.get('total') != len(selected) or len(rows) != len(selected) or ({cell_text(row.get('spuId')) for row in rows} != set(selected)):
            raise ValueError('explicit SPU scope is incomplete or unavailable to authenticated ERP')
        total = len(selected)
    elif target:
        payload = client.target_inventory_page(target)
        rows = payload['rows']
        if len(rows) != 1 or cell_text(rows[0].get('spuId')) != target or payload['total'] != 1:
            raise ValueError('target SPU is not uniquely available to authenticated ERP')
        total = 1
    elif max_spus is not None:
        if max_spus > 100:
            raise ValueError('diagnostic sample is limited to 100 SPUs')
        payload = client.inventory_page(1, max_spus)
        rows, total = (payload['rows'], payload['total'])
        if len(rows) != min(max_spus, total):
            raise ValueError('diagnostic product page is incomplete')
    else:
        rows = collect_self_operated_inventory(client)
        total = len(rows)
    if not rows:
        raise ValueError('authorized product source is empty; verify login and scope')
    empty_sku_spus = set()
    if callable(getattr(client, 'short_title_rows', None)):
        zero_sku_ids = [cell_text(row['spuId']) for row in rows if row.get('declaredSkuCount') == 0]
        for group in _chunks(zero_sku_ids, 5):
            observed = client.short_title_rows(group)
            if not isinstance(observed, list) or any((not isinstance(row, dict) or cell_text(row.get('productId')) not in group or (not re.fullmatch('\\d+', cell_text(row.get('skuId')))) for row in observed)):
                raise ValueError('zero-SKU verification returned an invalid product scope')
            nonempty = {cell_text(row['productId']) for row in observed}
            empty_sku_spus.update(set(group) - nonempty)
    result = _inspect_self_operated_rows(client, shop, rows, total=total, target=target, scope='selected' if selected else 'target' if target else 'sample' if max_spus is not None else 'full', inspect_batch_size=50, empty_sku_spus=empty_sku_spus)
    result.summary.update({'inventorySource': 'authorized-product-list', 'scopeMode': 'available-permissions'})
    if selected:
        result.summary['targetSpuIds'] = selected
    if callable(getattr(type(client), 'inspect_material_details', None)):
        result.summary['materialInspectionSource'] = 'exact-spu-detail'
    return _clear_self_operated_replacement_sources(result)

def _clear_self_operated_replacement_sources(result):
    for job in result.prepared.jobs:
        uploaded = result.existing_materials.get(job.spu_id, {}).get('uploaded_urls', {})
        replace_slots = set()
        if job.needs_scenes:
            replace_slots.update(('scene1', 'scene2'))
        if job.needs_selling_image:
            replace_slots.add('selling')
        if job.needs_white:
            replace_slots.add('transparent')
            if any((row.get('whiteStatus') in {'missing', 'rejected', 'invalid'} for row in result.report_rows if row['spuId'] == job.spu_id)):
                replace_slots.add('white')
        reports = [row for row in result.report_rows if row['spuId'] == job.spu_id]
        for slot in replace_slots:
            if not reports or any((row.get(f'{slot}Status') != 'approved' for row in reports)):
                uploaded.pop(slot, None)
    return result

def reconcile_self_operated_segment(original, jobs, inspections):
    return _protected_core.call('jd_material_agent.reconcile_self_operated_segment', locals())

def _inspect_self_operated_rows(client, shop: dict[str, Any], spu_rows: list[dict[str, Any]], *, total: int, target: str, scope: str, inspect_batch_size: int, empty_sku_spus: set[str] | None=None) -> SelfOperatedDiscoveryResult:
    inspection_by_spu: dict[str, dict[str, Any]] = {}
    empty_sku_spus = empty_sku_spus or set()
    requested_ids = [cell_text(item.get('spuId')) for item in spu_rows if cell_text(item.get('spuId')) not in empty_sku_spus]
    inspect = client.inspect_material_details if callable(getattr(type(client), 'inspect_material_details', None)) else client.inspect_spus
    for batch in _chunks(requested_ids, inspect_batch_size):
        results = inspect(batch)
        if not isinstance(results, list):
            raise ValueError('self-operated inspection returned an invalid batch')
        for result in results:
            spu_id = cell_text(result.get('spuId'))
            if spu_id not in batch or spu_id in inspection_by_spu:
                raise ValueError('self-operated inspection returned an unexpected or duplicate SPU')
            inspection_by_spu[spu_id] = result
        missing = set(batch) - set(inspection_by_spu)
        if missing:
            raise ValueError('self-operated inspection omitted requested SPUs')
    headers = ('item_sku_id', 'sku名称', 'spu_id', '短标是否维护', '卖点是否维护', '白底图是否维护', '场景图是否维护', '卖点图是否维护')
    prepared_rows: list[SourceRow] = []
    jobs: list[SpuJob] = []
    report_rows: list[dict[str, Any]] = []
    failed_spus: list[str] = []
    reference_urls: dict[str, str] = {}
    existing_materials: dict[str, dict[str, Any]] = {}
    for spu in spu_rows:
        spu_id = cell_text(spu.get('spuId'))
        if spu_id in empty_sku_spus:
            continue
        inspection = inspection_by_spu[spu_id]
        skus = inspection.get('skus') or []
        children = inspection.get('textChildren') or []
        child_by_sku = {cell_text(item.get('skuId')): item for item in children if cell_text(item.get('skuId'))}
        errors: list[str] = []
        if inspection.get('errors'):
            errors.append('readonly lookup incomplete')
        declared = int(spu.get('declaredSkuCount') or 0)
        sku_ids = [cell_text(item.get('skuId')) for item in skus]
        if not skus or any((not re.fullmatch('\\d+', sku_id) for sku_id in sku_ids)):
            errors.append('SKU expansion is empty or invalid')
        if len(set(sku_ids)) != len(sku_ids):
            errors.append('SKU expansion contains duplicates')
        if declared != len(skus):
            errors.append(f'declared SKU count {declared} does not match expanded count {len(skus)}')
        if set(child_by_sku) != set(sku_ids):
            errors.append('text detail SKU set does not match material SKU set')
        sku_reports: list[dict[str, Any]] = []
        slot_states: dict[str, list[str]] = {slot: [] for slot in SELF_OPERATED_IMAGE_SLOTS}
        slot_urls: dict[str, set[str]] = {slot: set() for slot in SELF_OPERATED_IMAGE_SLOTS}
        complete_point_sets: set[tuple[str, str, str]] = set()
        short_missing: list[str] = []
        for sku in skus:
            sku_id = cell_text(sku.get('skuId'))
            child = child_by_sku.get(sku_id, {})
            image_states = {}
            reused_urls = {}
            for slot, (material_type, order) in SELF_OPERATED_IMAGE_SLOTS.items():
                state, url = _self_image_status(sku, material_type, order)
                image_states[slot] = state
                slot_states[slot].append(state)
                if url:
                    slot_urls[slot].add(url)
                    reused_urls[slot] = url
                if state == 'unknown':
                    errors.append('material support information is unavailable')
            points = [cell_text(value) for value in child.get('sellPoints') or []]
            points = [value for value in points if value]
            points_status = 'complete' if len(points) >= 3 else 'missing'
            if points_status == 'complete':
                complete_point_sets.add(tuple(points[:3]))
            short_status, short_title = _self_short_title(sku, child)
            if short_status in {'unknown', 'conflict'}:
                errors.append('short title field is unavailable or inconsistent')
            elif short_status == 'missing':
                short_missing.append(sku_id)
            sku_reports.append({'shopName': cell_text(shop.get('shopName')), 'maskedShopId': cell_text(shop.get('maskedShopId')), 'spuId': spu_id, 'skuId': sku_id, 'productName': cell_text(spu.get('productName')) or cell_text(sku.get('skuName')), 'skuName': cell_text(sku.get('skuName')), 'score': cell_text(spu.get('score')), 'scoreStatus': _self_score_status(spu.get('score')), **{f'{slot}Status': image_states[slot] for slot in SELF_OPERATED_IMAGE_SLOTS}, 'sellingPointCount': len(points), 'sellingPointsStatus': points_status, 'shortTitle': short_title, 'shortTitleStatus': short_status, 'maintenanceItems': [], 'waitingItems': [slot for slot, state in image_states.items() if state == 'pending'], 'unsupportedItems': [slot for slot, state in image_states.items() if state == 'unsupported'], 'reusedMaterialUrls': reused_urls, 'inspectionError': '', 'materialTask': False, 'shortTitleTask': short_status == 'missing'})
            sku_reports[-1]['rejectedMaterials'] = [{'slot': slot, 'materialType': material_type, 'order': order, 'url': cell_text(material.get('url')), 'reason': cell_text(material.get('reason'))} for slot, (material_type, order) in SELF_OPERATED_IMAGE_SLOTS.items() for material in sku.get('materials') or [] if int(material.get('type') or 0) == material_type and int(material.get('order') or 0) == order and (int(material.get('status') or 0) == 5)]
        points_conflict = len(complete_point_sets) > 1
        if points_conflict:
            for row in sku_reports:
                if row['sellingPointsStatus'] == 'complete':
                    row['sellingPointsStatus'] = 'conflict'
        missing_slots: set[str] = set()
        waiting_slots: set[str] = set()
        for slot, states in slot_states.items():
            if 'pending' in states:
                waiting_slots.add(slot)
            elif any((state in {'missing', 'rejected', 'invalid'} for state in states)):
                missing_slots.add(slot)
        point_missing = points_conflict or any((row['sellingPointsStatus'] == 'missing' for row in sku_reports))
        needs_material = bool(missing_slots or point_missing)
        if errors:
            failed_spus.append(spu_id)
            message = '; '.join(dict.fromkeys(errors))
            for row in sku_reports:
                row['inspectionError'] = message
            report_rows.extend(sku_reports)
            continue
        uploaded_urls = {slot: next(iter(urls)) for slot, urls in slot_urls.items() if len(urls) == 1}
        existing: dict[str, Any] = {'uploaded_urls': uploaded_urls}
        if len(complete_point_sets) == 1:
            existing['selling_points'] = list(next(iter(complete_point_sets)))
        existing_materials[spu_id] = existing
        material_need_names = sorted(missing_slots | ({'points'} if point_missing else set()))
        score_status = _self_score_status(spu.get('score'), bool(material_need_names or short_missing))
        for row in sku_reports:
            row['scoreStatus'] = score_status
            row['maintenanceItems'] = material_need_names + (['shortTitle'] if row['shortTitleTask'] else [])
            row['waitingItems'] = sorted(set(row['waitingItems']) | waiting_slots)
            row['materialTask'] = needs_material
        if missing_slots:
            image = cell_text(spu.get('logo')) or cell_text(skus[0].get('logo'))
            try:
                reference_urls[spu_id] = _jd_image_url(image)
            except ValueError:
                failed_spus.append(spu_id)
                for row in sku_reports:
                    row['inspectionError'] = 'trusted reference image is unavailable'
                report_rows.extend(sku_reports)
                continue
        report_rows.extend(sku_reports)
        if not (needs_material or short_missing):
            continue
        title = cell_text(skus[0].get('skuName')) or cell_text(spu.get('productName'))
        raw_rows = []
        for sku, report in zip(skus, sku_reports):
            raw = (report['skuId'], report['skuName'], spu_id, 0 if report['shortTitleTask'] else 1, 0 if point_missing else 1, 0 if {'white', 'transparent'} & missing_slots else 1, 0 if {'scene1', 'scene2'} & missing_slots else 1, 0 if 'selling' in missing_slots else 1)
            raw_rows.append(SourceRow(report['skuId'], report['skuName'] or title, spu_id, cell_text(spu.get('score')), raw, '', '自动巡检底表'))
        prepared_rows.extend(raw_rows)
        jobs.append(SpuJob(spu_id=spu_id, representative_sku_id=raw_rows[0].sku_id, title=title, sku_ids=tuple((row.sku_id for row in raw_rows)), skus=tuple(((row.sku_id, row.title) for row in raw_rows)), short_title_sku_ids=tuple(short_missing), needs_selling_points=point_missing, needs_white=bool({'white', 'transparent'} & missing_slots), needs_scenes=bool({'scene1', 'scene2'} & missing_slots), needs_selling_image='selling' in missing_slots))
    failed_set = set(failed_spus)
    if failed_set:
        prepared_rows = [row for row in prepared_rows if row.spu_id not in failed_set]
        jobs = [job for job in jobs if job.spu_id not in failed_set]
        for spu_id in failed_set:
            reference_urls.pop(spu_id, None)
            existing_materials.pop(spu_id, None)
    prepared = PreparedSource('self-operated-live', '自动巡检底表', headers, tuple(prepared_rows), tuple(jobs), sum((len(inspection_by_spu[spu_id].get('skus') or []) for spu_id in requested_ids)), sum((len(inspection_by_spu[spu_id].get('skus') or []) for spu_id in requested_ids)) - len(prepared_rows), 0)
    summary = _summarize_self_operated_report(report_rows, total_spus=total, inspected_spus=len(spu_rows), failed_spus=len(failed_set), pagination_complete=len(spu_rows) == total)
    summary.update({'rejectedSpus': len({row['spuId'] for row in report_rows if row.get('rejectedMaterials')}), 'rejectedSkus': len({row['skuId'] for row in report_rows if row.get('rejectedMaterials')})})
    safe_to_continue = not failed_set and scope != 'sample'
    summary.update({'discoveryScope': scope, 'targetSpuId': target, 'fullInventoryVerified': scope == 'full' and summary['paginationComplete'], 'safeToContinue': safe_to_continue, 'emptySkuSpus': sorted(empty_sku_spus)})
    return SelfOperatedDiscoveryResult(prepared, reference_urls, existing_materials, tuple(report_rows), tuple(dict.fromkeys(failed_spus)), summary, cell_text(shop.get('shopName')), cell_text(shop.get('maskedShopId')), safe_to_continue)

def _summarize_self_operated_report(rows: list[dict[str, Any]], *, total_spus: int, inspected_spus: int, failed_spus: int, pagination_complete: bool) -> dict[str, Any]:
    slot_fields = {'white': 'whiteStatus', 'transparent': 'transparentStatus', 'scene1': 'scene1Status', 'scene2': 'scene2Status', 'selling': 'sellingStatus'}
    status_counts = {}
    for slot, field in slot_fields.items():
        status_counts[slot] = {state: sum((1 for row in rows if row.get(field) == state)) for state in ('approved', 'pending', 'rejected', 'missing', 'invalid', 'unsupported', 'unknown')}
    task_spus = {row['spuId'] for row in rows if row.get('materialTask') or row.get('shortTitleTask')}
    task_skus = {row['skuId'] for row in rows if row.get('materialTask') or row.get('shortTitleTask')}
    short_only = {row['skuId'] for row in rows if row.get('shortTitleTask') and (not row.get('materialTask'))}
    return {'onlineSpus': total_spus, 'inspectedSpus': inspected_spus, 'onlineSkusInScope': len(rows), 'lowScoreSpus': len({row['spuId'] for row in rows if row.get('scoreStatus') == 'low'}), 'lowScoreSkus': sum((1 for row in rows if row.get('scoreStatus') == 'low')), 'unscoredSpus': len({row['spuId'] for row in rows if row.get('scoreStatus') == 'unscored'}), 'unscoredSkus': sum((1 for row in rows if row.get('scoreStatus') == 'unscored')), 'imageStatusCounts': status_counts, 'sellingPointsMissingSkus': sum((1 for row in rows if row.get('sellingPointsStatus') == 'missing')), 'sellingPointsConflictSpus': len({row['spuId'] for row in rows if row.get('sellingPointsStatus') == 'conflict'}), 'sellingPointsConflictSkus': sum((1 for row in rows if row.get('sellingPointsStatus') == 'conflict')), 'shortTitleMissingSkus': sum((1 for row in rows if row.get('shortTitleStatus') == 'missing')), 'maintenanceSpus': len(task_spus), 'maintenanceSkus': len(task_skus), 'shortTitleOnlySkus': len(short_only), 'incompleteSpus': failed_spus, 'paginationComplete': pagination_complete, 'safeToContinue': failed_spus == 0}

def write_self_operated_discovery(result: SelfOperatedDiscoveryResult, output_dir: str | Path) -> dict[str, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    json_path = output / '自营商品信息分实时巡检.json'
    report_path = output / '自营商品信息分实时巡检.xlsx'
    prepared_path = output / '自营商品信息分自动底表.xlsx'
    payload = {'skillVersion': SKILL_VERSION, 'shopName': result.shop_name, 'maskedShopId': result.masked_shop_id, 'summary': result.summary, 'failedSpus': list(result.failed_spus), 'rows': list(result.report_rows)}
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    workbook = Workbook()
    detail = workbook.active
    detail.title = '巡检明细'
    report_headers = ('店铺名称', '店铺ID(脱敏)', 'SPUID', 'SKUID', '商品名称', 'SKU名称', '商品信息分', '分数状态', '白底图状态', '透明图状态', '场景图1状态', '场景图2状态', '卖点图状态', '商品卖点数量', '商品卖点状态', '短标题', '短标题状态', '待维护项目', '审核中项目', '不支持项目', '已复用素材URL', '巡检错误', '进入素材任务', '进入短标题任务')
    detail.append(report_headers)
    for row in result.report_rows:
        detail.append([row['shopName'], row['maskedShopId'], row['spuId'], row['skuId'], row['productName'], row['skuName'], row['score'], row['scoreStatus'], row['whiteStatus'], row['transparentStatus'], row['scene1Status'], row['scene2Status'], row['sellingStatus'], row['sellingPointCount'], row['sellingPointsStatus'], row['shortTitle'], row['shortTitleStatus'], ','.join(row['maintenanceItems']), ','.join(row['waitingItems']), ','.join(row['unsupportedItems']), json.dumps(row['reusedMaterialUrls'], ensure_ascii=False), row['inspectionError'], 1 if row['materialTask'] else 0, 1 if row['shortTitleTask'] else 0])
    summary_sheet = workbook.create_sheet('巡检汇总')
    summary_sheet.append(['指标', '值'])
    for key, value in result.summary.items():
        summary_sheet.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list, tuple)) else value])
    workbook.save(report_path)
    workbook.close()
    prepared_book = Workbook()
    prepared_sheet = prepared_book.active
    prepared_sheet.title = '自动巡检底表'
    prepared_sheet.append(list(result.prepared.headers))
    for row in result.prepared.rows:
        prepared_sheet.append(list(row.raw))
    prepared_book.save(prepared_path)
    prepared_book.close()
    return {'json': json_path, 'reportWorkbook': report_path, 'preparedSource': prepared_path}

def split_rows_by_spu(rows: Iterable[SourceRow], maximum_spus: int=MAX_SPUS_PER_BATCH) -> list[list[SourceRow]]:
    if maximum_spus < 1:
        raise ValueError('maximum_spus 必须大于 0。')
    groups: dict[str, list[SourceRow]] = {}
    for row in rows:
        groups.setdefault(row.spu_id, []).append(row)
    chunks, current = ([], [])
    for index, items in enumerate(groups.values()):
        if index and index % maximum_spus == 0:
            chunks.append(current)
            current = []
        current.extend(items)
    if current:
        chunks.append(current)
    return chunks

def selling_point_prompts(title: str) -> tuple[str, str]:
    return _protected_core.call('jd_material_agent.selling_point_prompts', locals())

def short_title_prompts(skus: Iterable[tuple[str, str]]) -> tuple[str, str]:
    return _protected_core.call('jd_material_agent.short_title_prompts', locals())

def _shared_image_rules(title: str | None) -> str:
    return _protected_core.call('jd_material_agent._shared_image_rules', locals())

def image_prompt(kind: str, title: str | None) -> str:
    return _protected_core.call('jd_material_agent.image_prompt', locals())

def _json_object(text: str) -> dict:
    cleaned = re.sub('^```(?:json)?\\s*|\\s*```$', '', cell_text(text), flags=re.I)
    start, end = (cleaned.find('{'), cleaned.rfind('}'))
    if start < 0 or end < start:
        raise ValueError('模型未返回 JSON 对象。')
    return json.loads(cleaned[start:end + 1])

class RequestGate:

    def __init__(self, requests_per_minute: int | None=None, start_interval: float=0.0):
        self.requests_per_minute = requests_per_minute
        self.start_interval = start_interval
        self.lock = threading.Lock()
        self.starts = deque()
        self.last_start = 0.0

    def wait(self) -> None:
        while True:
            with self.lock:
                now = time.monotonic()
                while self.starts and now - self.starts[0] >= 60:
                    self.starts.popleft()
                delay = max(0.0, self.start_interval - (now - self.last_start))
                if self.requests_per_minute and len(self.starts) >= self.requests_per_minute:
                    delay = max(delay, 60 - (now - self.starts[0]))
                if delay <= 0:
                    self.starts.append(now)
                    self.last_start = now
                    return
            time.sleep(min(delay, 1.0))

class ModelQuotaExceeded(RuntimeError):

    def __init__(self):
        super().__init__('模型接口明确返回当前模型配额已用尽（code=2007）；已暂停新模型请求，恢复额度后使用原目录继续。')

class ModelTransientError(RuntimeError):
    pass

class ImageModelResourceBlocked(RuntimeError):

    def __init__(self, evidence=None):
        self.evidence = evidence
        super().__init__('image model resource temporarily blocked by provider content policy; stop image requests and contact the provider; do not rotate keys to bypass')

def is_image_request_url(url):
    return urlparse(url).path.endswith(('/images/edits', '/images/gemini_flash/generations'))

def _model_resource_failure_evidence(response, api_key):
    sensitive = {'authorization', 'cookie', 'setcookie', 'apikey', 'token', 'accesstoken', 'refreshtoken', 'idtoken', 'secret', 'clientsecret', 'password', 'b64json', 'base64', 'image', 'images', 'imagedata'}

    def redact(value):
        if isinstance(value, dict):
            return {str(name).replace(api_key, '[REDACTED]'): '[REDACTED]' if re.sub('[^a-z0-9]', '', str(name).lower()) in sensitive else redact(item) for name, item in value.items()}
        if isinstance(value, list):
            return [redact(item) for item in value]
        if not isinstance(value, str):
            return value
        clean = value.replace(api_key, '[REDACTED]')
        try:
            nested = json.loads(clean)
        except ValueError:
            nested = None
        if isinstance(nested, (dict, list)):
            return json.dumps(redact(nested), ensure_ascii=False)
        clean = re.sub('data:image/[^;\\s]+;base64,[A-Za-z0-9+/=]+', '[REDACTED_IMAGE]', clean)
        clean = re.sub('(?im)^(\\s*(?:authorization|cookie|set-cookie)\\s*:\\s*).*$', '\\1[REDACTED]', clean)
        clean = re.sub('(?i)\\bBearer\\s+[^\\s"\',;]+', 'Bearer [REDACTED]', clean)
        return re.sub('(?i)\\b(api[_-]?key|access[_-]?token|refresh[_-]?token|token|password|secret)\\s*[=:]\\s*[^\\r\\n,;}]+', '\\1=[REDACTED]', clean)
    try:
        request = getattr(response, '_request', None)
        if request is not None and urlparse(str(request.url)).path.endswith('/images/gemini_flash/generations'):
            from gemini_image_protocol import response_summary
            body = response_summary(response)
        else:
            body = response.json()
    except ValueError:
        body = response.text
    allowed = {'request-id', 'x-request-id', 'x-ms-request-id', 'apim-request-id', 'x-trace-id', 'traceparent', 'x-correlation-id', 'date', 'retry-after'}
    return {'httpStatus': response.status_code, 'responseHeaders': redact({name: value for name, value in response.headers.items() if name.lower() in allowed}), 'response': redact(body)}

class ModelClient:

    def __init__(self, api_key: str, timeout: float=600.0, image_concurrency: int | None=None):
        if not api_key or not api_key.isascii() or any((character.isspace() for character in api_key)):
            raise ValueError(f'{PROVIDER.api_key_env} 未设置或格式无效。')
        self.http = httpx.Client(timeout=timeout, headers={'Authorization': f'Bearer {api_key}'})
        self._api_key = api_key
        self.text_gate = RequestGate(PROVIDER.requests_per_minute, PROVIDER.text_start_interval_seconds)
        self.image_gate = self.text_gate if PROVIDER.requests_per_minute else RequestGate(None, PROVIDER.image_start_interval_seconds)
        self.text_semaphore = threading.Semaphore(PROVIDER.text_concurrency)
        requested = image_concurrency or PROVIDER.default_image_concurrency
        self.image_semaphore = threading.Semaphore(max(1, min(PROVIDER.maximum_image_concurrency, requested)))
        self.quota_exhausted = False
        self.authentication_failed = False
        self.image_resource_blocked = False
        self.image_resource_failure = None
        self.cancel_event = None

    def close(self) -> None:
        self.http.close()

    def _request(self, method: str, url: str, *, request_gate=None, max_attempts=4, **kwargs) -> httpx.Response:
        last_error = None
        if getattr(self, 'no_generation_retries', False):
            max_attempts = 1
        image_request = is_image_request_url(url)
        recovery = getattr(self, 'image_resource_recovery', None) if image_request else None
        if recovery is not None:
            cached = recovery.before_request(method, url, kwargs)
            if cached is not None:
                return cached
        for attempt in range(max_attempts):
            if image_request and self.image_resource_blocked:
                raise ImageModelResourceBlocked(self.image_resource_failure)
            if self.authentication_failed:
                raise RuntimeError('model authentication failed; further dispatch stopped')
            if self.cancel_event is not None and self.cancel_event.is_set():
                raise RuntimeError('generation pipeline stopped before starting another request')
            if self.quota_exhausted:
                raise ModelQuotaExceeded()
            if request_gate is not None:
                request_gate.wait()
            if self.cancel_event is not None and self.cancel_event.is_set():
                raise RuntimeError('generation pipeline stopped before dispatch')
            if self.quota_exhausted:
                raise ModelQuotaExceeded()
            try:
                if image_request and self.image_resource_blocked:
                    raise ImageModelResourceBlocked(self.image_resource_failure)
                if self.authentication_failed:
                    raise RuntimeError('model authentication failed; further dispatch stopped')
                if image_request and callable(getattr(self, 'before_image_dispatch', None)):
                    self.before_image_dispatch(method, url, kwargs)
                response = self.http.request(method, url, **kwargs)
                if image_request and callable(getattr(self, 'observe_image_response', None)):
                    self.observe_image_response(response)
                if response.status_code < 400:
                    return response
                if recovery is not None:
                    recovery.last_response_evidence = _model_resource_failure_evidence(response, self._api_key)
                if image_request and all((marker in response.text.lower() for marker in ('forbidden', 'resource', 'temporarily blocked', 'content policy'))):
                    self.image_resource_blocked = True
                    evidence = _model_resource_failure_evidence(response, self._api_key)
                    if self.image_resource_failure is None:
                        self.image_resource_failure = evidence
                    if recovery is not None:
                        recovery.record_failure(method, url, kwargs, evidence)
                    raise ImageModelResourceBlocked(evidence)
                if response.status_code in (401, 403):
                    self.authentication_failed = True
                    raise RuntimeError('model authentication failed; further dispatch stopped')
                if response.status_code == 429:
                    try:
                        payload = response.json()
                    except ValueError:
                        payload = None
                    details = payload.get('error') if isinstance(payload, dict) else None
                    if isinstance(details, dict) and str(details.get('code')) == '2007':
                        self.quota_exhausted = True
                        raise ModelQuotaExceeded()
                error_body = response.text[:2000]
                if urlparse(url).path.endswith('/images/gemini_flash/generations'):
                    error_body = json.dumps(_model_resource_failure_evidence(response, self._api_key), ensure_ascii=False)
                error = RuntimeError(f'HTTP {response.status_code}: {error_body}')
                if getattr(self, 'skip_failed_images', False) and response.status_code == 400 and image_request and is_moderation_error(error):
                    raise error
                retryable_gateway_error = response.status_code == 400 and ('BAD_UPSTREAM' in response.text.upper() or 'FAILED_RESPONSE' in response.text.upper() or bool(re.search('\\\\?"code\\\\?"\\s*:\\s*(?:500|502|503|504)', response.text, re.IGNORECASE)))
                if response.status_code not in {408, 409, 425, 429} and response.status_code < 500 and (not retryable_gateway_error):
                    raise error
                last_error = error
                retry_after = response.headers.get('Retry-After')
                if response.status_code == 429:
                    try:
                        delay = float(retry_after) if retry_after else 60.0
                        if delay <= 0:
                            delay = 60.0
                    except ValueError:
                        delay = 60.0
                else:
                    delay = min(60, 2 ** attempt)
            except (httpx.TimeoutException, httpx.TransportError) as error:
                last_error = error
                delay = min(60, 2 ** attempt)
            self.last_request_retry_delay = delay
            if attempt < max_attempts - 1:
                time.sleep(delay)
        raise ModelTransientError(str(last_error) or '模型请求失败。') from last_error

    def text(self, system: str, user: str) -> str:
        with self.text_semaphore:
            url = f'{PROVIDER.base_url}/{PROVIDER.text_endpoint}'
            if PROVIDER.id == 'jd-internal':
                payload = {'model': PROVIDER.text_model, 'instructions': system, 'input': user, 'stream': False}
            else:
                payload = {'model': PROVIDER.text_model, 'messages': [{'role': 'system', 'content': system}, {'role': 'user', 'content': user}]}
            value = self._request('POST', url, json=payload, request_gate=self.text_gate).json()
        if PROVIDER.id == 'jd-internal':
            if isinstance(value.get('output_text'), str) and value['output_text'].strip():
                return value['output_text'].strip()
            parts = []
            for item in value.get('output') or []:
                for content in item.get('content') or []:
                    text = content.get('text') or content.get('value')
                    if isinstance(text, str):
                        parts.append(text)
            if parts:
                return '\n'.join(parts)
        else:
            return cell_text(value.get('choices', [{}])[0].get('message', {}).get('content'))
        raise ValueError('文本响应中没有内容。')

    def edit_image(self, prompt: str, source_path: Path) -> bytes:
        with self.image_semaphore:
            url = f'{PROVIDER.base_url}/{PROVIDER.image_endpoint}'
            source_bytes = source_path.read_bytes()
            with Image.open(BytesIO(source_bytes)) as image:
                mime = image.get_format_mimetype() or 'image/png'
            if PROVIDER.id == 'jd-internal':
                data_url = f"data:{mime};base64,{base64.b64encode(source_bytes).decode('ascii')}"
                payload = {'model': PROVIDER.image_model, 'prompt': prompt, 'size': '1024x1024', 'response_format': 'b64_json', 'image': [data_url]}
                value = self._request('POST', url, json=payload, request_gate=self.image_gate).json()
            else:
                files = {'image': (f'reference{source_path.suffix.lower()}', source_bytes, mime)}
                data = {'model': PROVIDER.image_model, 'prompt': prompt, 'size': '1024x1024', 'output_format': 'png'}
                value = self._request('POST', url, data=data, files=files, request_gate=self.image_gate).json()
            item = (value.get('data') or [{}])[0]
            encoded = item.get('b64_json') or item.get('b64Json')
            if not encoded and isinstance(item.get('url'), str) and item['url'].startswith('data:image/'):
                encoded = item['url'].split(',', 1)[1]
            if not encoded:
                raise ValueError('图片响应中没有 base64 数据。')
            return base64.b64decode(encoded, validate=True)

    def selling_points(self, title: str) -> list[str]:
        system, user = selling_point_prompts(title)
        last_error = None
        attempts = 1 if getattr(self, 'no_generation_retries', False) else 3
        for _ in range(attempts):
            try:
                return validate_selling_points(_json_object(self.text(system, user))['selling_points'])
            except (KeyError, TypeError, ValueError) as error:
                last_error = error
                user += '\n上次输出不合规，请只重写 JSON 并严格遵守长度和禁用词。'
        raise ValueError(f'连续{attempts}次未生成合规卖点：{last_error}')

    def short_titles(self, skus: Iterable[tuple[str, str]]) -> dict[str, str]:
        skus = tuple(skus)
        expected = {sku_id for sku_id, _ in skus}
        system, user = short_title_prompts(skus)
        last_error = None
        attempts = 1 if getattr(self, 'no_generation_retries', False) else 3
        for _ in range(attempts):
            try:
                rows = _json_object(self.text(system, user))['short_titles']
                result = {cell_text(item['sku_id']): cell_text(item['short_title']) for item in rows}
                if set(result) != expected or any((not valid_short_title(value) for value in result.values())):
                    raise ValueError('短标题存在缺失、额外 SKU 或长度不合规。')
                return result
            except (KeyError, TypeError, ValueError) as error:
                last_error = error
                user += '\n上次输出不合规，请只重写缺失或长度不合规项并保持 SKUID 精确一致。'
        raise ValueError(f'连续{attempts}次未生成合规短标题：{last_error}')

def _square_800(image: Image.Image, background: tuple[int, ...]) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    image.thumbnail((800, 800), Image.Resampling.LANCZOS)
    canvas = Image.new(image.mode, (800, 800), background)
    offset = ((800 - image.width) // 2, (800 - image.height) // 2)
    canvas.paste(image, offset, image if image.mode == 'RGBA' else None)
    return canvas

def normalize_jpeg(content: bytes, output: Path, pure_white: bool=False) -> None:
    with Image.open(BytesIO(content)) as opened:
        image = _square_800(opened.convert('RGB'), (255, 255, 255))
    if pure_white:
        corners = ((0, 0), (799, 0), (0, 799), (799, 799))
        if any((min(image.getpixel(point)) < 248 for point in corners)):
            inset = image.copy()
            inset.thumbnail((760, 760), Image.Resampling.LANCZOS)
            image = Image.new('RGB', (800, 800), (255, 255, 255))
            image.paste(inset, ((800 - inset.width) // 2, (800 - inset.height) // 2))
    output.parent.mkdir(parents=True, exist_ok=True)
    for quality in (95, 92, 88, 84, 80, 75):
        image.save(output, 'JPEG', quality=quality, optimize=True, progressive=True, subsampling=0)
        if pure_white and output.stat().st_size <= 1000000 and (not valid_material_image(output, 'white')):
            inset = image.copy()
            inset.thumbnail((760, 760), Image.Resampling.LANCZOS)
            image = Image.new('RGB', (800, 800), (255, 255, 255))
            image.paste(inset, ((800 - inset.width) // 2, (800 - inset.height) // 2))
            image.save(output, 'JPEG', quality=quality, optimize=True, progressive=True, subsampling=0)
        if output.stat().st_size <= 1000000:
            return
    raise ValueError(f'JPG 超过 1MB：{output}')

def make_transparent(source: Path, output: Path) -> None:
    with Image.open(source) as opened:
        rgb = _square_800(opened.convert('RGB'), (255, 255, 255))
    width, height = rgb.size
    pixels = rgb.load()
    visited = bytearray(width * height)
    background = bytearray(width * height)
    queue = deque()

    def enqueue(x: int, y: int) -> None:
        index = y * width + x
        if not visited[index]:
            visited[index] = 1
            queue.append((x, y))
    for x in range(width):
        enqueue(x, 0)
        enqueue(x, height - 1)
    for y in range(height):
        enqueue(0, y)
        enqueue(width - 1, y)
    while queue:
        x, y = queue.popleft()
        pixel = pixels[x, y]
        if min(pixel) < 248 or max(pixel) - min(pixel) > 8:
            continue
        background[y * width + x] = 1
        if x:
            enqueue(x - 1, y)
        if x + 1 < width:
            enqueue(x + 1, y)
        if y:
            enqueue(x, y - 1)
        if y + 1 < height:
            enqueue(x, y + 1)
    rgba = rgb.convert('RGBA')
    alpha = Image.new('L', (width, height), 255)
    alpha_pixels = alpha.load()
    for y in range(height):
        for x in range(width):
            if background[y * width + x]:
                alpha_pixels[x, y] = 0
    rgba.putalpha(alpha)
    output.parent.mkdir(parents=True, exist_ok=True)
    rgba.save(output, 'PNG', optimize=True)
    if output.stat().st_size > 3000000:
        raise ValueError(f'透明 PNG 超过 3MB：{output}')

def valid_material_image(path: Path, kind: str) -> bool:
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            if image.size != (800, 800):
                return False
            if kind == 'transparent':
                return image.format == 'PNG' and 'A' in image.mode and (path.stat().st_size <= 3000000)
            if image.format not in {'JPEG', 'JPG'} or path.stat().st_size > 1000000:
                return False
            if kind == 'white':
                rgb = image.convert('RGB')
                return all((min(rgb.getpixel(point)) >= 248 for point in ((0, 0), (799, 0), (0, 799), (799, 799))))
            return True
    except (OSError, ValueError):
        return False

def resolve_reference(reference_dir: Path, job: SpuJob) -> Path:
    candidates = []
    for stem in (job.spu_id, f'{job.spu_id}_{job.representative_sku_id}', job.representative_sku_id):
        for suffix in ('.png', '.jpg', '.jpeg', '.webp'):
            candidates.append(reference_dir / f'{stem}{suffix}')
    match = next((path for path in candidates if path.is_file()), None)
    if not match:
        raise FileNotFoundError(f'SPU {job.spu_id} 缺少参考图。请由 Agent 使用已登录京东页面下载主图到 {reference_dir}，命名为 {job.spu_id}.jpg。')
    return match

def gallery_reference_candidates(reference_dir: Path, job: SpuJob) -> list[Path]:
    pattern = re.compile(f'^{re.escape(job.spu_id)}_gallery_(\\d+)$', re.I)
    candidates = []
    for path in reference_dir.iterdir() if reference_dir.is_dir() else ():
        if path.suffix.lower() not in {'.png', '.jpg', '.jpeg', '.webp'}:
            continue
        match = pattern.fullmatch(path.stem)
        if match:
            candidates.append((int(match.group(1)), path.name.lower(), path))
    return [item[2] for item in sorted(candidates)]

def is_moderation_error(error: BaseException) -> bool:
    return bool(re.search('moderation|content[_ -]?policy|safety\\s+system|审核|安全策略', str(error), re.I))

def _trusted_jd_image_url(url: str) -> bool:
    parsed = urlparse(url)
    host = (parsed.hostname or '').lower()
    return parsed.scheme == 'https' and (host == '360buyimg.com' or host.endswith('.360buyimg.com'))

def restore_uploaded_white(url: str, output: Path, timeout: float) -> None:
    if not _trusted_jd_image_url(url):
        raise ValueError('已确认白底图不是可信京东图片域名，拒绝下载。')
    response = httpx.get(url, timeout=timeout, follow_redirects=True)
    response.raise_for_status()
    if len(response.content) > 20000000:
        raise ValueError('已确认白底图超过 20MB，拒绝处理。')
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(response.content)
    if not valid_material_image(output, 'white'):
        normalize_jpeg(response.content, output, pure_white=True)

def _request_image_set(client: ModelClient, kinds: Iterable[str], source: Path, title: str | None) -> tuple[dict[str, bytes], dict[str, Exception]]:
    kinds = tuple(kinds)
    results: dict[str, bytes] = {}
    errors: dict[str, Exception] = {}
    if not kinds:
        return (results, errors)
    with ThreadPoolExecutor(max_workers=min(len(kinds), PROVIDER.maximum_image_concurrency)) as image_pool:
        routed = getattr(client, 'edit_material', None)
        futures = {image_pool.submit(routed, kind, image_prompt(kind, title), source) if callable(routed) else image_pool.submit(client.edit_image, image_prompt(kind, title), source): kind for kind in kinds}
        for future in as_completed(futures):
            kind = futures[future]
            try:
                results[kind] = future.result()
            except _protected_core.ProtectedCoreError:
                raise
            except Exception as error:
                errors[kind] = error
    return (results, errors)

def _commit_staged_images(results: dict[str, bytes], images: Path, names: dict[str, str], spu_id: str) -> None:
    staging = images.parent / '.state' / 'staging' / spu_id
    shutil.rmtree(staging, ignore_errors=True)
    staging.mkdir(parents=True, exist_ok=True)
    try:
        for kind, content in results.items():
            normalize_jpeg(content, staging / names[kind], pure_white=kind == 'white')
        for kind in results:
            (staging / names[kind]).replace(images / names[kind])
    finally:
        shutil.rmtree(staging, ignore_errors=True)

def _load_state(path: Path) -> dict:
    if not path.exists():
        return {'version': 1, 'jobs': {}}
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        return {'version': 1, 'jobs': {}}

def _save_state(path: Path, state: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)

def _load_upload_ledger(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding='utf-8'))
        if not isinstance(payload, dict) or not isinstance(payload.get('files', {}), dict):
            raise ValueError('invalid upload ledger')
        payload.setdefault('version', 1)
        payload.setdefault('files', {})
        return payload
    except FileNotFoundError:
        return {'version': 1, 'files': {}}

def _save_upload_ledger(path: Path, ledger: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(ledger, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)

def _write_url_export(path: Path, prepared: PreparedSource, state: dict) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '图片空间链接'
    sheet.append(['图片名称', '图片链接'])
    for job in prepared.jobs:
        uploaded = (state.get('jobs', {}).get(job.spu_id, {}) or {}).get('uploaded_urls') or {}
        for kind, name in image_names(job.spu_id).items():
            if uploaded.get(kind):
                sheet.append([name, uploaded[kind]])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path

def _normalize_upload_filename(value: object) -> str:
    return re.sub('\\s+', '', cell_text(value)).casefold()

def _self_operated_remote_image(item: dict[str, Any]) -> dict[str, str]:
    name = cell_text(item.get('name') or item.get('imgName'))
    image_id = cell_text(item.get('imageId') or item.get('imgId') or item.get('id'))
    raw_path = item.get('path') or item.get('imgUrl') or item.get('url')
    url = _jd_image_url(raw_path)
    path = _jfs_path(raw_path)
    if not name or not image_id or (not url) or (not path):
        raise ValueError('image-space response is missing name, image ID, or jfs path')
    return {'imageId': image_id, 'name': name, 'path': path, 'url': url}

def _scan_self_operated_image_space(client: Any, category_id: int, *, page_size: int=50, query: str='', first_page=None) -> list[dict[str, str]]:
    if type(page_size) is not int or page_size < 1:
        raise ValueError('self-operated image-space page size must be a positive integer')
    rows: list[dict[str, str]] = []
    identities: set[str] = set()
    expected_pages: int | None = None
    expected_total: int | None = None
    page = 1
    while True:
        if page == 1 and first_page is not None:
            payload = first_page
        elif query:
            payload = client.image_page(int(category_id), page, page_size, query)
        else:
            payload = client.image_page(int(category_id), page, page_size)
        returned_page = int(payload.get('page') or page)
        page_total = int(payload.get('pageTotal') or 0)
        total = int(payload.get('total') or 0)
        page_rows = payload.get('images') or []
        if returned_page != page:
            raise ValueError('image-space pagination returned an unexpected page')
        if expected_pages is None:
            expected_pages = page_total
            expected_total = total
        elif page_total != expected_pages or total != expected_total:
            raise ValueError('image-space pagination changed during the scan')
        if page < max(expected_pages, 1) and (not page_rows):
            raise ValueError('image-space pagination returned an empty intermediate page')
        for raw in page_rows:
            normalized = _self_operated_remote_image(raw)
            identity = normalized['imageId'] or normalized['path']
            if identity in identities:
                raise ValueError('image-space pagination returned a duplicate image')
            identities.add(identity)
            rows.append(normalized)
        if page >= max(expected_pages, 1):
            break
        page += 1
    return rows

def _scan_self_operated_image_queries(client, category_id, queries):
    queries = list(dict.fromkeys(queries))
    rows = []
    batched = callable(getattr(type(client), 'image_pages', None))
    for group in _chunks(queries, 10):
        pages = {}
        if batched:
            replies = client.image_pages(int(category_id), group, page_size=50)
            if not isinstance(replies, list) or len(replies) != len(group) or any((not isinstance(item, dict) or not isinstance(item.get('query'), str) or (not isinstance(item.get('payload'), dict)) for item in replies)) or ({item['query'] for item in replies} != set(group)):
                raise ValueError('image lookup response must cover the exact requested queries')
            pages = {item['query']: item['payload'] for item in replies}
        for query in group:
            rows.extend(_scan_self_operated_image_space(client, int(category_id), query=query, first_page=pages.get(query)))
    return rows

def _validate_uploaded_image_url(url: str, timeout: float=20.0, attempts: int=5, *, expected_sha256: str='') -> None:
    trusted = _jd_image_url(url)
    if not _jfs_path(trusted):
        raise ValueError('uploaded image URL does not contain a jfs path')
    last_error: Exception | None = None
    for attempt in range(attempts):
        try:
            response = httpx.get(trusted, timeout=timeout, follow_redirects=True)
        except _protected_core.ProtectedCoreError:
            raise
        except Exception as error:
            last_error = error
        else:
            content_type = response.headers.get('content-type', '')
            if response.status_code == 200 and content_type.lower().startswith('image/') and response.content:
                if expected_sha256 and hashlib.sha256(response.content).hexdigest() != expected_sha256:
                    raise ValueError('uploaded image content hash does not match the approved file')
                return
            last_error = ValueError(f'uploaded image GET verification failed: HTTP {response.status_code} {content_type}')
        if attempt + 1 < attempts:
            time.sleep(1)
    raise ValueError(str(last_error or 'uploaded image GET verification failed'))

def upload_self_operated_materials(client: Any, prepared: PreparedSource, output_dir: str | Path, category_id: int=0, *, images_dir: str | Path | None=None, selected_spu_ids: set[str] | None=None) -> UploadResult:
    import upload_quarantine
    upload_quarantine.assert_write_targets(output_dir, selected_spu_ids if selected_spu_ids is not None else {job.spu_id for job in prepared.jobs})
    images = Path(images_dir) if images_dir else Path(output_dir) / '全部图片'
    files = []
    for job in prepared.jobs:
        if selected_spu_ids is not None and job.spu_id not in selected_spu_ids:
            continue
        required = set(_model_kinds(job)) | ({'transparent'} if job.needs_white else set())
        files.extend((images / name for kind, name in image_names(job.spu_id).items() if kind in required and (images / name).is_file()))
    staging = client.stage_upload_files(files) if callable(getattr(type(client), 'stage_upload_files', None)) else nullcontext()
    with staging:
        return _upload_self_operated_materials(client, prepared, output_dir, category_id, images_dir=images_dir, selected_spu_ids=selected_spu_ids)

def _submit_self_operated_upload_batch(client, pending, category_id, ledger, ledger_path, submitted, failures):
    try:
        results = client.upload_images([item['path'] for item in pending], int(category_id), expected_hashes={item['path'].name: ledger['files'][item['filename']]['sha256'] for item in pending})
        expected = {item['remoteName'] for item in pending}
        if not isinstance(results, list) or len(results) != len(pending) or any((not isinstance(item, dict) or item.get('status') not in {'submitted', 'unconfirmed', 'not-submitted'} for item in results)) or ({item.get('fileName') for item in results} != expected):
            raise ValueError('batch upload response does not cover the exact submitted file set')
        by_name = {item['fileName']: item for item in results}
    except _protected_core.ProtectedCoreError:
        raise
    except Exception as error:
        for item in pending:
            ledger['files'][item['filename']]['status'] = 'unconfirmed'
            failures[item['filename']] = f"{item['filename']}: {error}"
        _save_upload_ledger(ledger_path, ledger)
        return
    for item in pending:
        result = by_name[item['remoteName']]
        entry = ledger['files'][item['filename']]
        if result['status'] != 'submitted':
            entry['status'] = result['status']
            failures[item['filename']] = f"{item['filename']}: batch upload {result['status']}"
        else:
            try:
                response = _self_operated_remote_image(result.get('response') or {})
                if response['name'] != item['remoteName']:
                    raise ValueError('image-space API returned a different filename')
                entry.update({'path': response['path'], 'url': response['url'], 'imageId': response['imageId'], 'status': 'submitted', 'submittedAt': int(time.time())})
                submitted.append((item['spuId'], item['kind'], item['filename'], response))
            except _protected_core.ProtectedCoreError:
                raise
            except Exception as error:
                entry['status'] = 'unconfirmed'
                failures[item['filename']] = f"{item['filename']}: {error}"
        _save_upload_ledger(ledger_path, ledger)

def _upload_self_operated_materials(client: Any, prepared: PreparedSource, output_dir: str | Path, category_id: int=0, *, images_dir: str | Path | None=None, selected_spu_ids: set[str] | None=None) -> UploadResult:
    output = Path(output_dir)
    transport = 'self-operated-direct-http-sff' if getattr(client, 'image_transport', 'browser') == 'direct-http' else 'self-operated-browser-sff'
    images = Path(images_dir) if images_dir else output / '全部图片'
    state_path = output / '.state' / 'task.json'
    state = _load_state(state_path)
    _bind_state(state, prepared, state_path)
    ledger_path = output / '.state' / 'upload-ledger.json'
    ledger = _load_upload_ledger(ledger_path)
    export_path = output / '图片空间链接.xlsx'
    selected_jobs = [job for job in prepared.jobs if selected_spu_ids is None or job.spu_id in selected_spu_ids]
    required_filenames = []
    query_prefixes = {}
    for job in selected_jobs:
        required = set(_model_kinds(job))
        if job.needs_white:
            required.add('transparent')
        names = image_names(job.spu_id)
        uploaded_urls = state.get('jobs', {}).get(job.spu_id, {}).get('uploaded_urls') or {}
        required_filenames.extend((names[kind] for kind in IMAGE_SUFFIXES if kind in required and (not uploaded_urls.get(kind))))
        query_prefixes.update({names[kind]: job.spu_id for kind in IMAGE_SUFFIXES if kind in required})
    query_names = []
    for filename in dict.fromkeys(required_filenames):
        entry = ledger['files'].get(filename) or {}
        remote_name = cell_text(entry.get('remoteName')) or filename
        query_names.append(query_prefixes[filename] if remote_name == filename else remote_name)
    remote_rows = _scan_self_operated_image_queries(client, int(category_id), query_names)
    remote_by_name: dict[str, list[dict[str, str]]] = {}
    for row in remote_rows:
        remote_by_name.setdefault(_normalize_upload_filename(row['name']), []).append(row)
    uploaded_count = 0
    reused_count = 0
    failures_by_name: dict[str, str] = {}
    submitted: list[tuple[str, str, str, dict[str, str]]] = []
    batch_supported = callable(getattr(type(client), 'upload_images', None)) and bool(getattr(client, 'expected_erp', True))
    batch_target = getattr(client, 'upload_batch_size', 5) if batch_supported else 5
    if type(batch_target) is not int or batch_target < 1:
        raise ValueError('upload batch target must be a positive integer')
    pending_uploads = []
    for job in selected_jobs:
        state_item = state['jobs'].setdefault(job.spu_id, {})
        uploaded_urls = dict(state_item.get('uploaded_urls') or {})
        required = set(_model_kinds(job))
        if job.needs_white:
            required.add('transparent')
        names = image_names(job.spu_id)
        for kind in IMAGE_SUFFIXES:
            if kind not in required:
                continue
            path = images / names[kind]
            digest = hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else ''
            entry = ledger['files'].get(path.name)
            remote_name = cell_text((entry or {}).get('remoteName')) or path.name
            if entry and digest and (entry.get('sha256') != digest):
                raise ValueError(f'uploaded filename content changed: {path.name}')
            if uploaded_urls.get(kind):
                reused_count += 1
                continue
            exact_rows = remote_by_name.get(_normalize_upload_filename(remote_name), [])
            if entry and digest and (entry.get('sha256') == digest):
                expected_path = cell_text(entry.get('path'))
                matched = next((row for row in exact_rows if row['path'] == expected_path), None)
                if not expected_path and len(exact_rows) == 1:
                    try:
                        _validate_uploaded_image_url(exact_rows[0]['url'], expected_sha256=digest)
                        matched = exact_rows[0]
                    except _protected_core.ProtectedCoreError:
                        raise
                    except Exception as error:
                        failures_by_name[path.name] = f'{path.name}: unconfirmed upload content could not be reconciled: {error}'
                        continue
                if matched:
                    _validate_uploaded_image_url(matched['url'])
                    entry.update({'status': 'verified', 'url': matched['url'], 'imageId': matched['imageId'], 'path': matched['path']})
                    uploaded_urls[kind] = matched['url']
                    state_item['uploaded_urls'] = uploaded_urls
                    _save_upload_ledger(ledger_path, ledger)
                    _save_state(state_path, state)
                    reused_count += 1
                    continue
                if entry.get('status') in {'uploading', 'unconfirmed', 'submitted', 'verified'}:
                    failures_by_name[path.name] = f'{path.name}: previous upload outcome requires exact reconciliation; not resubmitted'
                    continue
            if not valid_material_image(path, kind):
                failures_by_name[path.name] = f'{path.name}: local image is missing or invalid'
                continue
            if exact_rows:
                remote_name = f'{path.stem}__{digest[:16]}{path.suffix}'
                alias_rows = remote_by_name.get(_normalize_upload_filename(remote_name), [])
                if alias_rows:
                    try:
                        if len(alias_rows) != 1:
                            raise ValueError('content-addressed filename is ambiguous')
                        matched = alias_rows[0]
                        _validate_uploaded_image_url(matched['url'], expected_sha256=digest)
                        ledger['files'][path.name] = {**matched, 'remoteName': remote_name, 'sha256': digest, 'kind': kind, 'status': 'verified', 'transport': 'verified-existing-hash'}
                        uploaded_urls[kind] = matched['url']
                        state_item['uploaded_urls'] = uploaded_urls
                        _save_upload_ledger(ledger_path, ledger)
                        _save_state(state_path, state)
                        reused_count += 1
                    except _protected_core.ProtectedCoreError:
                        raise
                    except Exception as error:
                        failures_by_name[path.name] = f'{path.name}: {error}'
                    continue
            upload_path = path
            if remote_name != path.name:
                if Path(remote_name).name != remote_name or '/' in remote_name or '\\' in remote_name:
                    raise ValueError('remote upload filename is not a safe basename')
                upload_path = output / '.state/upload-files' / remote_name
                upload_path.parent.mkdir(parents=True, exist_ok=True)
                if not upload_path.exists():
                    shutil.copyfile(path, upload_path)
                if hashlib.sha256(upload_path.read_bytes()).hexdigest() != digest:
                    raise ValueError('upload alias content differs from the approved image')
            ledger['files'][path.name] = {'sha256': digest, 'kind': kind, 'remoteName': remote_name, 'status': 'uploading', 'submittedAt': int(time.time()), 'transport': transport}
            _save_upload_ledger(ledger_path, ledger)
            if batch_supported:
                pending_uploads.append({'spuId': job.spu_id, 'kind': kind, 'filename': path.name, 'path': upload_path, 'remoteName': remote_name})
                if len(pending_uploads) >= batch_target:
                    _submit_self_operated_upload_batch(client, pending_uploads, category_id, ledger, ledger_path, submitted, failures_by_name)
                    pending_uploads = []
                continue
            try:
                response = _self_operated_remote_image(client.upload_image(upload_path, int(category_id)))
                if response['name'] != remote_name:
                    raise ValueError('image-space API returned a different filename')
                ledger['files'][path.name] = {'sha256': digest, 'kind': kind, 'path': response['path'], 'url': response['url'], 'imageId': response['imageId'], 'remoteName': response['name'], 'status': 'submitted', 'submittedAt': int(time.time()), 'transport': transport}
                _save_upload_ledger(ledger_path, ledger)
                submitted.append((job.spu_id, kind, path.name, response))
            except _protected_core.ProtectedCoreError:
                raise
            except Exception as error:
                ledger['files'][path.name]['status'] = 'unconfirmed'
                _save_upload_ledger(ledger_path, ledger)
                failures_by_name[path.name] = f'{path.name}: {error}'
    if pending_uploads:
        _submit_self_operated_upload_batch(client, pending_uploads, category_id, ledger, ledger_path, submitted, failures_by_name)
    if submitted:
        verified_rows = []
        query_groups = {}
        for spu_id, kind, filename, response in submitted:
            query_groups.setdefault(spu_id, set()).add(_normalize_upload_filename(response['name']))
        for query, expected_names in query_groups.items():
            for attempt in range(5):
                rows = _scan_self_operated_image_space(client, int(category_id), query=query)
                returned_names = {_normalize_upload_filename(row['name']) for row in rows}
                if expected_names.issubset(returned_names) or attempt == 4:
                    verified_rows.extend(rows)
                    break
                time.sleep(1)
        verified_by_name: dict[str, list[dict[str, str]]] = {}
        for row in verified_rows:
            verified_by_name.setdefault(_normalize_upload_filename(row['name']), []).append(row)
        candidates = []
        for spu_id, kind, filename, response in submitted:
            matches = [row for row in verified_by_name.get(_normalize_upload_filename(response['name']), []) if row['imageId'] == response['imageId'] and row['path'] == response['path']]
            if len(matches) != 1:
                failures_by_name[filename] = f'{filename}: exact upload readback did not match image ID and jfs path'
                continue
            candidates.append((spu_id, kind, filename, matches[0]['url']))
        if candidates:
            with ThreadPoolExecutor(max_workers=4) as executor:
                futures = {}
                try:
                    for item in candidates:
                        futures[executor.submit(_validate_uploaded_image_url, item[3])] = item
                    for future in as_completed(futures):
                        spu_id, kind, filename, url = futures[future]
                        try:
                            future.result()
                        except _protected_core.ProtectedCoreError:
                            raise
                        except Exception as error:
                            failures_by_name[filename] = f'{filename}: {error}'
                            continue
                        entry = ledger['files'][filename]
                        entry.update({'status': 'verified', 'url': url, 'verifiedAt': int(time.time())})
                        state_item = state['jobs'].setdefault(spu_id, {})
                        uploaded_urls = dict(state_item.get('uploaded_urls') or {})
                        uploaded_urls[kind] = url
                        state_item['uploaded_urls'] = uploaded_urls
                        _save_upload_ledger(ledger_path, ledger)
                        _save_state(state_path, state)
                        uploaded_count += 1
                finally:
                    for future in futures:
                        future.cancel()
    _write_url_export(export_path, prepared, state)
    failures = tuple(failures_by_name.values())
    failed_spu_ids = tuple((job.spu_id for job in selected_jobs if any((name in failures_by_name for name in image_names(job.spu_id).values()))))
    return UploadResult(uploaded_count, reused_count, len(failures), False, export_path, failures, failed_spu_ids)

def _business_mode(prepared: PreparedSource) -> str:
    if VARIANT_ID == 'merchant':
        return 'merchant'
    return 'self-operated' if prepared.source_type in {'self-operated', 'self-operated-live'} else 'pop'

def _task_signature(prepared: PreparedSource) -> str:
    payload = {'source_type': prepared.source_type, 'jobs': [{'spu_id': job.spu_id, 'skus': list(job.skus), 'short_title_sku_ids': list(job.short_title_sku_ids), 'needs_selling_points': job.needs_selling_points, 'needs_white': job.needs_white, 'needs_scenes': job.needs_scenes, 'needs_selling_image': job.needs_selling_image} for job in prepared.jobs]}
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()

def _prepared_payload(prepared: PreparedSource) -> dict[str, Any]:
    return {'sourceType': prepared.source_type, 'sheetName': prepared.sheet_name, 'headers': list(prepared.headers), 'rows': [asdict(row) for row in prepared.rows], 'jobs': [asdict(job) for job in prepared.jobs], 'originalRows': prepared.original_rows, 'filteredRows': prepared.filtered_rows, 'duplicateRows': prepared.duplicate_rows}

def _prepared_from_payload(payload: dict[str, Any]) -> PreparedSource:
    rows = tuple((SourceRow(sku_id=cell_text(item.get('sku_id')), title=cell_text(item.get('title')), spu_id=cell_text(item.get('spu_id')), score=cell_text(item.get('score')), raw=tuple(item.get('raw') or ()), source_path=cell_text(item.get('source_path')), sheet_name=cell_text(item.get('sheet_name')), row_number=int(item.get('row_number') or 0), header_row=int(item.get('header_row') or 1)) for item in payload.get('rows') or []))
    jobs = tuple((SpuJob(spu_id=cell_text(item.get('spu_id')), representative_sku_id=cell_text(item.get('representative_sku_id')), title=cell_text(item.get('title')), sku_ids=tuple((cell_text(value) for value in item.get('sku_ids') or ())), skus=tuple(((cell_text(pair[0]), cell_text(pair[1])) for pair in item.get('skus') or ())), short_title_sku_ids=tuple((cell_text(value) for value in item.get('short_title_sku_ids') or ())), needs_selling_points=bool(item.get('needs_selling_points')), needs_white=bool(item.get('needs_white')), needs_scenes=bool(item.get('needs_scenes')), needs_selling_image=bool(item.get('needs_selling_image'))) for item in payload.get('jobs') or []))
    return PreparedSource(source_type=cell_text(payload.get('sourceType')), sheet_name=cell_text(payload.get('sheetName')), headers=tuple(payload.get('headers') or ()), rows=rows, jobs=jobs, original_rows=int(payload.get('originalRows') or 0), filtered_rows=int(payload.get('filteredRows') or 0), duplicate_rows=int(payload.get('duplicateRows') or 0))

def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()

def _auto_maintain_plan_hash(payload: dict[str, Any]) -> str:
    content = {key: value for key, value in payload.items() if key != 'planHash'}
    encoded = json.dumps(content, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()

def _bind_state(state: dict, prepared: PreparedSource, state_path: Path) -> None:
    mode = _business_mode(prepared)
    signature = _task_signature(prepared)
    existing_jobs = set((state.get('jobs') or {}).keys())
    current_jobs = {job.spu_id for job in prepared.jobs}
    existing_variant = state.get('variant')
    existing_mode = state.get('business_mode')
    existing_signature = state.get('task_signature')
    if existing_variant and existing_variant != VARIANT_ID:
        raise ValueError('输出目录属于其他 Skill，禁止串用任务状态。请选择新的输出目录。')
    if prepared.source_type == 'self-operated-live' and existing_mode == 'pop' and (existing_signature == signature):
        existing_mode = 'self-operated'
    if existing_mode and existing_mode != mode:
        raise ValueError('POP 与自营模式隔离：当前输出目录属于另一模式，请选择新的输出目录。')
    if existing_signature and existing_signature != signature:
        raise ValueError('输出目录属于另一份底表任务，禁止复用其状态。请选择新的输出目录。')
    if not existing_signature and existing_jobs and (not existing_jobs.issubset(current_jobs)):
        raise ValueError('旧任务状态与当前底表不一致，禁止串用缓存。请选择新的输出目录。')
    state['version'] = 2
    state['variant'] = VARIANT_ID
    state['business_mode'] = mode
    state['task_signature'] = signature
    state.setdefault('jobs', {})
    _save_state(state_path, state)

def _model_kinds(job: SpuJob) -> list[str]:
    kinds = []
    if job.needs_white:
        kinds.append('white')
    if job.needs_scenes:
        kinds.extend(('scene1', 'scene2'))
    if job.needs_selling_image:
        kinds.append('selling')
    return kinds

def _has_material_work(job: SpuJob) -> bool:
    return job.needs_selling_points or bool(_model_kinds(job))

def _material_is_complete(job: SpuJob, item: dict, images: Path) -> bool:
    if job.needs_selling_points:
        try:
            validate_selling_points(item.get('selling_points') or [])
        except ValueError:
            return False
    names = image_names(job.spu_id)
    required = _model_kinds(job)
    if job.needs_white:
        required.append('transparent')
    uploaded = item.get('uploaded_urls') or {}
    return all((valid_material_image(images / names[kind], kind) or bool(uploaded.get(kind)) for kind in required))

def _short_titles_are_complete(job: SpuJob, item: dict) -> bool:
    values = item.get('short_titles') or {}
    return all((valid_short_title(values.get(sku_id)) for sku_id in job.short_title_sku_ids))

def _write_results(prepared: PreparedSource, output: Path, state: dict, failures: list[tuple[SpuJob, str]]) -> BatchResult:
    output.mkdir(parents=True, exist_ok=True)
    images = output / '全部图片'
    reported_failures: dict[str, list[str]] = {}
    for job, reason in failures:
        reported_failures.setdefault(job.spu_id, []).append(cell_text(reason) or '处理失败。')
    completed = []
    material_successes = []
    incomplete: list[tuple[SpuJob, str]] = []
    for job in prepared.jobs:
        item = state.get('jobs', {}).get(job.spu_id, {})
        material_complete = not _has_material_work(job) or _material_is_complete(job, item, images)
        short_complete = _short_titles_are_complete(job, item)
        reasons = list(reported_failures.get(job.spu_id, []))
        if not material_complete:
            reasons.append('计划素材未全部完成。')
        if not short_complete:
            reasons.append('计划短标题未全部完成。')
        if _has_material_work(job) and material_complete:
            material_successes.append(job)
        if material_complete and short_complete and (not reasons):
            completed.append(job)
        else:
            incomplete.append((job, '；'.join(dict.fromkeys(reasons))))
    manifest_path = output / '素材生成清单.xlsx' if material_successes else None
    short_title_path = output / '商品导入短标题模板_已回填.xlsx'
    failure_path = output / '失败SPU清单.xlsx' if incomplete else None
    if manifest_path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '素材生成清单'
        sheet.append(['SPUID', '代表SKUID', 'SKU数量', '代表商品名称', '卖点1', '卖点2', '卖点3', '白底图', '透明图', '场景图1', '场景图2', '卖点图', '需要卖点', '需要白底图', '需要场景图', '需要卖点图'])
        for job in material_successes:
            item = state['jobs'].get(job.spu_id, {})
            points = item.get('selling_points') if job.needs_selling_points else ['', '', '']
            names = image_names(job.spu_id)
            sheet.append([job.spu_id, job.representative_sku_id, job.sku_count, job.title, *points, names['white'] if job.needs_white else '', names['transparent'] if job.needs_white else '', names['scene1'] if job.needs_scenes else '', names['scene2'] if job.needs_scenes else '', names['selling'] if job.needs_selling_image else '', int(job.needs_selling_points), int(job.needs_white), int(job.needs_scenes), int(job.needs_selling_image)])
        detail = workbook.create_sheet('SKU明细')
        detail.append(['SPUID', 'SKUID', '商品名称'])
        for job in material_successes:
            for sku_id, title in job.skus:
                detail.append([job.spu_id, sku_id, title])
        workbook.save(manifest_path)
        workbook.close()
    else:
        (output / '素材生成清单.xlsx').unlink(missing_ok=True)
    short_rows = []
    for job in prepared.jobs:
        values = state['jobs'].get(job.spu_id, {}).get('short_titles', {})
        for sku_id in job.short_title_sku_ids:
            if valid_short_title(values.get(sku_id)):
                short_rows.append((sku_id, values[sku_id]))
    if short_rows:
        from manual_import_export import TITLE_TEMPLATE, TITLE_MAX_BYTES, write_parts
        snapshot = output / '.state/title-workbooks' / _auto_maintain_plan_hash({'rows': short_rows})
        parts = write_parts(short_rows, TITLE_TEMPLATE, '批量维护短标题', snapshot, '批量维护短标题', max_bytes=TITLE_MAX_BYTES)
        short_title_path = Path(parts[0]['path'])
    else:
        short_title_path = None
        (output / '商品导入短标题模板_已回填.xlsx').unlink(missing_ok=True)
    if failure_path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '失败SPU'
        sheet.append(['SPUID', 'SKUID', '商品名称', '失败原因'])
        for job, reason in incomplete:
            for sku_id, title in job.skus:
                sheet.append([job.spu_id, sku_id, title, reason])
        workbook.save(failure_path)
        workbook.close()
    else:
        (output / '失败SPU清单.xlsx').unlink(missing_ok=True)
    return BatchResult(len(prepared.jobs), len(completed), len(incomplete), output, manifest_path, short_title_path, failure_path)

def generate_batch(prepared: PreparedSource, reference_dir: str | Path, output_dir: str | Path, *, spu_concurrency: int=5, image_concurrency: int | None=None, timeout: float=600.0, existing_url_exports: Iterable[str | Path]=(), existing_materials: dict[str, dict[str, Any]] | None=None, model_client_factory=None) -> BatchResult:
    import upload_quarantine
    quarantined = upload_quarantine.for_batch(output_dir)
    jobs = tuple((job for job in prepared.jobs if job.spu_id not in quarantined))
    if not 1 <= spu_concurrency <= MAX_SPU_CONCURRENCY:
        raise ValueError('SPU 并发必须为 1-50。')
    output = Path(output_dir)
    images = output / '全部图片'
    state_path = output / '.state' / 'task.json'
    images.mkdir(parents=True, exist_ok=True)
    reference_dir = Path(reference_dir)
    state = _load_state(state_path)
    _bind_state(state, prepared, state_path)
    for job in jobs:
        seed = (existing_materials or {}).get(job.spu_id) or {}
        if not seed:
            continue
        item = state['jobs'].setdefault(job.spu_id, {})
        uploaded = dict(item.get('uploaded_urls') or {})
        for kind, url in (seed.get('uploaded_urls') or {}).items():
            uploaded.setdefault(kind, url)
        if uploaded:
            item['uploaded_urls'] = uploaded
        if job.needs_selling_points and (not item.get('selling_points')):
            points = seed.get('selling_points') or []
            if len(points) == 3:
                item['selling_points'] = list(points)
    _save_state(state_path, state)
    if existing_url_exports:
        existing_urls = load_url_exports(existing_url_exports)
        for job in jobs:
            names = image_names(job.spu_id)
            uploaded = dict(state['jobs'].setdefault(job.spu_id, {}).get('uploaded_urls') or {})
            for kind, filename in names.items():
                url = existing_urls.get(normalize_material_stem(filename))
                if url:
                    uploaded[kind] = url
            if uploaded:
                state['jobs'][job.spu_id]['uploaded_urls'] = uploaded
        _save_state(state_path, state)
    state_lock = threading.Lock()

    def needs_model_client() -> bool:
        for job in jobs:
            item = state.get('jobs', {}).get(job.spu_id, {})
            if job.needs_selling_points and (not item.get('selling_points')):
                return True
            if not _short_titles_are_complete(job, item):
                return True
            names = image_names(job.spu_id)
            uploaded = item.get('uploaded_urls') or {}
            if any((not valid_material_image(images / names[kind], kind) and (not uploaded.get(kind)) for kind in _model_kinds(job))):
                return True
        return False
    client = None
    if needs_model_client():
        client = model_client_factory() if model_client_factory else ModelClient(os.environ.get(PROVIDER.api_key_env, ''), timeout=timeout, image_concurrency=image_concurrency)
    failures: list[tuple[SpuJob, str]] = [(job, upload_quarantine.REASON) for job in prepared.jobs if job.spu_id in quarantined]

    def save_job(spu_id: str, key: str, value) -> None:
        with state_lock:
            state['jobs'].setdefault(spu_id, {})[key] = value
            _save_state(state_path, state)

    def complete_job(spu_id: str) -> None:
        with state_lock:
            item = state['jobs'].setdefault(spu_id, {})
            item['status'] = 'completed'
            item.pop('failure', None)
            _save_state(state_path, state)

    def process_text(job: SpuJob, item: dict) -> None:
        if job.needs_selling_points and (not item.get('selling_points')):
            assert client is not None
            save_job(job.spu_id, 'selling_points', client.selling_points(job.title))
        pending_short = [(sku, title) for sku, title in job.skus if sku in job.short_title_sku_ids and (not valid_short_title(item.get('short_titles', {}).get(sku)))]
        if pending_short:
            assert client is not None
            merged = dict(item.get('short_titles', {}))
            for index in range(0, len(pending_short), 50):
                merged.update(client.short_titles(pending_short[index:index + 50]))
                save_job(job.spu_id, 'short_titles', merged)

    def process_images(job: SpuJob, item: dict) -> None:
        names = image_names(job.spu_id)
        uploaded = item.get('uploaded_urls') or {}
        required_model_kinds = _model_kinds(job)
        pending = [kind for kind in required_model_kinds if not valid_material_image(images / names[kind], kind) and (not uploaded.get(kind))]
        if pending:
            assert client is not None
            recovery = dict(item.get('moderation_recovery') or {})
            skip_failed = getattr(client, 'skip_failed_images', False) is True
            if skip_failed and (item.get('skipped_image_generation') or recovery.get('normal_status') == 'blocked'):
                raise RuntimeError('previous image generation failed; skipped without regeneration')
            source = resolve_reference(reference_dir, job)
            if recovery.get('normal_status') != 'blocked':
                results, errors = _request_image_set(client, pending, source, job.title)
                for kind, content in results.items():
                    normalize_jpeg(content, images / names[kind], pure_white=kind == 'white')
                if any((isinstance(error, ImageModelResourceBlocked) for error in errors.values())):
                    raise ImageModelResourceBlocked()
                cancellation = getattr(client, 'cancel_event', None)
                if cancellation is not None and cancellation.is_set() is True:
                    raise RuntimeError('generation cancelled; partial cache preserved without skipping images')
                if errors and skip_failed:
                    if not getattr(client, 'quota_exhausted', False) and (not getattr(client, 'authentication_failed', False)):
                        save_job(job.spu_id, 'skipped_image_generation', {kind: str(error) for kind, error in errors.items()})
                    raise RuntimeError('；'.join((f'{kind}: {error}' for kind, error in errors.items())))
                all_required_moderated = set(pending) == set(required_model_kinds) and len(errors) == len(required_model_kinds) and all((is_moderation_error(error) for error in errors.values()))
                if errors and (not all_required_moderated):
                    raise RuntimeError('；'.join((f'{kind}: {error}' for kind, error in errors.items())))
                if all_required_moderated:
                    recovery['normal_status'] = 'blocked'
                    save_job(job.spu_id, 'moderation_recovery', recovery)
            pending = [kind for kind in required_model_kinds if not valid_material_image(images / names[kind], kind)]
            if pending and recovery.get('normal_status') == 'blocked':
                if recovery.get('titleless_status') not in {'failed', 'succeeded'}:
                    recovery['titleless_status'] = 'running'
                    save_job(job.spu_id, 'moderation_recovery', recovery)
                    results, errors = _request_image_set(client, required_model_kinds, source, None)
                    if not errors and set(results) == set(required_model_kinds):
                        _commit_staged_images(results, images, names, job.spu_id)
                        recovery['titleless_status'] = 'succeeded'
                        recovery['status'] = 'succeeded'
                        save_job(job.spu_id, 'moderation_recovery', recovery)
                    else:
                        recovery['titleless_status'] = 'failed'
                        save_job(job.spu_id, 'moderation_recovery', recovery)
                        if not errors or not all((is_moderation_error(error) for error in errors.values())):
                            raise RuntimeError('；'.join((f'{kind}: {error}' for kind, error in errors.items())))
                pending = [kind for kind in required_model_kinds if not valid_material_image(images / names[kind], kind)]
                if pending:
                    attempted = set(recovery.get('attempted_candidates') or [])
                    available = [path for path in gallery_reference_candidates(reference_dir, job) if path.name not in attempted]
                    selected = None
                    for candidate in available:
                        attempted.add(candidate.name)
                        recovery['attempted_candidates'] = sorted(attempted)
                        recovery['status'] = 'running'
                        save_job(job.spu_id, 'moderation_recovery', recovery)
                        try:
                            white = client.edit_image(image_prompt('white', None), candidate)
                        except _protected_core.ProtectedCoreError:
                            raise
                        except Exception as error:
                            if is_moderation_error(error):
                                continue
                            recovery['status'] = 'failed'
                            save_job(job.spu_id, 'moderation_recovery', recovery)
                            raise
                        normalize_jpeg(white, images / names['white'], pure_white=True)
                        if job.needs_white:
                            make_transparent(images / names['white'], images / names['transparent'])
                        derived = [kind for kind in required_model_kinds if kind != 'white' and (not valid_material_image(images / names[kind], kind))]
                        results, errors = _request_image_set(client, derived, images / names['white'], None)
                        if errors:
                            recovery['status'] = 'failed'
                            save_job(job.spu_id, 'moderation_recovery', recovery)
                            raise RuntimeError('；'.join((f'{kind}: {error}' for kind, error in errors.items())))
                        _commit_staged_images(results, images, names, job.spu_id)
                        selected = candidate
                        break
                    if not selected:
                        recovery['status'] = 'failed'
                        save_job(job.spu_id, 'moderation_recovery', recovery)
                        raise RuntimeError(f'SPU {job.spu_id} 正常轮和无标题轮均被安全策略拦截；请由 Agent 获取其他可信商品相册图，按 {job.spu_id}_gallery_001.jpg 顺序放入参考图目录后重试。')
                    recovery['selected_candidate'] = selected.name
                    recovery['status'] = 'succeeded'
                    save_job(job.spu_id, 'moderation_recovery', recovery)
        if job.needs_white and (not uploaded.get('transparent')) and (not valid_material_image(images / names['transparent'], 'transparent')):
            if not valid_material_image(images / names['white'], 'white'):
                if uploaded.get('white'):
                    restore_uploaded_white(uploaded['white'], images / names['white'], timeout)
                else:
                    raise RuntimeError('缺少有效白底图，无法生成透明图。')
            make_transparent(images / names['white'], images / names['transparent'])

    def process(job: SpuJob) -> None:
        with state_lock:
            item = dict(state['jobs'].setdefault(job.spu_id, {}))
        errors = []
        with ThreadPoolExecutor(max_workers=2) as branches:
            futures = [branches.submit(process_text, job, item), branches.submit(process_images, job, item)]
            for future in futures:
                try:
                    future.result()
                except _protected_core.ProtectedCoreError:
                    raise
                except Exception as error:
                    errors.append(str(error))
        if errors:
            raise RuntimeError('；'.join(dict.fromkeys(errors)))
        complete_job(job.spu_id)
    try:
        with ThreadPoolExecutor(max_workers=min(spu_concurrency, max(1, len(jobs)))) as executor:
            futures = {executor.submit(process, job): job for job in jobs}
            for future in as_completed(futures):
                job = futures[future]
                try:
                    future.result()
                except _protected_core.ProtectedCoreError:
                    raise
                except Exception as error:
                    failures.append((job, str(error)))
                    save_job(job.spu_id, 'failure', str(error))
    finally:
        if client and model_client_factory is None:
            client.close()
    result = _write_results(prepared, output, state, failures)
    if getattr(client, 'image_resource_blocked', False) is True:
        state['model_pause'] = {'reason': 'image-resource-blocked', 'providerResolutionRequired': True}
        evidence = getattr(client, 'image_resource_failure', None)
        if evidence is not None:
            state['model_pause']['evidence'] = evidence
            evidence_path = state_path.parent / 'model-failures' / f'{time.time_ns()}.json'
            _save_state(evidence_path, evidence)
            state['model_pause']['evidencePath'] = str(evidence_path.relative_to(output))
            state['model_pause']['evidenceSha256'] = _file_sha256(evidence_path)
        _save_state(state_path, state)
        raise ImageModelResourceBlocked(evidence)
    if getattr(client, 'quota_exhausted', False) is True:
        state['model_pause'] = {'reason': 'quota-exhausted', 'code': 2007}
        _save_state(state_path, state)
        raise ModelQuotaExceeded()
    if getattr(client, 'authentication_failed', False) is True:
        state['model_pause'] = {'reason': 'authentication-failed'}
        _save_state(state_path, state)
        raise RuntimeError('model authentication failed; cached progress preserved')
    cancellation = getattr(client, 'cancel_event', None)
    if cancellation is not None and cancellation.is_set() is True:
        state['model_pause'] = {'reason': 'generation-cancelled'}
        _save_state(state_path, state)
        raise RuntimeError('generation cancelled; partial cache preserved without skipping images')
    if state.pop('model_pause', None) is not None:
        _save_state(state_path, state)
    return result

def normalize_material_stem(value: object) -> str:
    name = cell_text(value).replace('\\', '/').rsplit('/', 1)[-1]
    return re.sub('\\.(?:jpe?g|png|webp)$', '', name, flags=re.I).strip()

def load_url_exports(paths: Iterable[str | Path]) -> dict[str, str]:
    mapping = {}
    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            found = False
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                for index, row in enumerate(rows[:20]):
                    headers = [cell_text(value).replace(' ', '') for value in row]
                    name_index = next((i for i, value in enumerate(headers) if value in {'图片名称', '文件名称', '名称'}), None)
                    url_index = next((i for i, value in enumerate(headers) if value in {'图片链接', '图片URL', '链接', 'URL'}), None)
                    if name_index is None or url_index is None:
                        continue
                    found = True
                    for data in rows[index + 1:]:
                        name = normalize_material_stem(data[name_index] if name_index < len(data) else None)
                        url = cell_text(data[url_index] if url_index < len(data) else None)
                        if not name and (not url):
                            continue
                        if url.startswith('//'):
                            url = 'https:' + url
                        if not re.match('^https?://', url):
                            raise ValueError(f'图片链接无效：{url}')
                        if name in mapping and mapping[name] != url:
                            raise ValueError(f'图片名称重复且链接不同：{name}')
                        mapping[name] = url
                    break
            if not found:
                raise ValueError(f'URL 表 {path} 未找到图片名称/图片链接表头。')
        finally:
            workbook.close()
    return mapping

def _download_reference(url: str, output: Path, timeout: float) -> Path:
    trusted_url = _jd_image_url(url)
    if output.is_file():
        try:
            with Image.open(output) as image:
                image.verify()
            return output
        except OSError:
            pass
    response = httpx.get(trusted_url, timeout=timeout, follow_redirects=True)
    response.raise_for_status()
    final_url = _jd_image_url(str(response.url))
    if not final_url:
        raise ValueError('product image download redirected to an invalid URL')
    if len(response.content) > 20 * 1024 * 1024:
        raise ValueError('product image exceeds 20MB')
    with Image.open(BytesIO(response.content)) as opened:
        image = ImageOps.exif_transpose(opened).convert('RGB')
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_suffix('.tmp.jpg')
        image.save(temporary, 'JPEG', quality=95, optimize=True)
        temporary.replace(output)
    return output

def _subset_prepared(prepared: PreparedSource, rows: list[SourceRow]) -> PreparedSource:
    selected_spus = list(dict.fromkeys((row.spu_id for row in rows)))
    jobs_by_spu = {job.spu_id: job for job in prepared.jobs}
    missing = [spu_id for spu_id in selected_spus if spu_id not in jobs_by_spu]
    if missing:
        raise ValueError(f"prepared subset contains unknown SPUs: {','.join(missing)}")
    selected_jobs = tuple((job for job in prepared.jobs if job.spu_id in selected_spus))
    rows_by_spu = {spu_id: [] for spu_id in selected_spus}
    for row in rows:
        rows_by_spu[row.spu_id].append(row.sku_id)
    if any((len(rows_by_spu[job.spu_id]) != len(job.sku_ids) or set(rows_by_spu[job.spu_id]) != set(job.sku_ids) for job in selected_jobs)):
        raise ValueError('prepared subset must retain every sibling SKU exactly once')
    return PreparedSource(prepared.source_type, prepared.sheet_name, prepared.headers, tuple(rows), selected_jobs, len(rows), 0, 0)

def _generate_erp_materials(result, args, output, *, model_client_factory=None):
    import upload_quarantine
    for index, rows in enumerate(split_rows_by_spu(result.prepared.rows, maximum_spus=50), 1):
        batch = _subset_prepared(result.prepared, rows)
        folder = output / f'批次{index:03d}'
        references = folder / '参考图'
        quarantined = upload_quarantine.for_batch(folder)
        for job in batch.jobs:
            if job.spu_id in quarantined or not _model_kinds(job):
                continue
            path = references / f'{job.spu_id}.jpg'
            if path.is_file():
                continue
            url = result.reference_urls.get(job.spu_id)
            if url:
                try:
                    _download_reference(url, path, args.timeout)
                except _protected_core.ProtectedCoreError:
                    raise
                except Exception:
                    pass
        kwargs = {'model_client_factory': model_client_factory} if model_client_factory else {}
        generate_batch(batch, references, folder, spu_concurrency=args.spu_concurrency, image_concurrency=args.image_concurrency, timeout=args.timeout, existing_materials={job.spu_id: result.existing_materials.get(job.spu_id, {}) for job in batch.jobs}, **kwargs)

@contextmanager
def _erp_execution_lock(output):
    lock = output / '.state' / 'execution.lock'
    lock.parent.mkdir(parents=True, exist_ok=True)
    try:
        descriptor = os.open(str(lock), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as error:
        raise ValueError('self-operated execution is active or interrupted; verify the recorded process before clearing its lock') from error
    try:
        with os.fdopen(descriptor, 'w', encoding='utf-8') as handle:
            handle.write(str(os.getpid()))
        yield
    finally:
        lock.unlink(missing_ok=True)
if __name__ == '__main__':
    from self_operated_cli import main
    raise SystemExit(main())
