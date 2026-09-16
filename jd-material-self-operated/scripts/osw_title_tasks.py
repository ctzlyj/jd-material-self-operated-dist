import base64
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import time
from openpyxl import load_workbook
import manual_import_export as export

def save(path, value):
    import first_use
    first_use.core._save_state(path, value)

def read(path):
    return json.loads(path.read_text(encoding='utf-8'))

def sha(data):
    return hashlib.sha256(data).hexdigest()

def parse_log(data, rows):
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        if len(book.sheetnames) != 1:
            raise ValueError('OSW log worksheet count changed')
        values = list(book.active.values)
    finally:
        book.close()
    expected = {row['skuId']: row for row in rows}
    if not values or list(values[0][:2]) != ['skuId', '短标题'] or len(values) != len(rows) + 1:
        raise ValueError('OSW row log does not cover the exact request')
    outcomes = []
    seen = set()
    for value in values[1:]:
        if len(value) != 3:
            raise ValueError('OSW row log schema changed')
        sku, title, status = map(lambda item: str(item or ''), value)
        if sku not in expected or sku in seen or title != expected[sku]['shortTitle']:
            raise ValueError('OSW row log identity or title changed')
        if status != '成功' and (not status.startswith('失败')):
            raise ValueError('OSW row outcome is not definitive')
        seen.add(sku)
        reason = status.split('{', 1)[0] if status != '成功' else ''
        outcomes.append({'skuId': sku, 'successKnown': True, 'success': status == '成功', 'message': reason, 'code': 'OSW_ROW_REJECTED' if status != '成功' else ''})
    return outcomes

def checked_task(reply, name, erp, task_id=None):
    rows = reply.get('rows')
    if not isinstance(rows, list) or reply.get('total') != len(rows) or len(rows) != 1:
        raise ValueError('OSW task not uniquely reconciled; never recreate')
    task = rows[0]
    if task.get('name') != name or task.get('erp') != erp or task.get('type') != 272 or (not re.fullmatch('[0-9a-f]{24}', str(task.get('id', '')))) or (task_id and task['id'] != task_id):
        raise ValueError('OSW task identity does not match original request')
    return task

def prepare_route(output):
    path = output / '.state/short-title-write.json'
    if not path.exists():
        return
    previous = read(path)
    if previous.get('transport') == 'osw':
        return
    if previous.get('unconfirmedTitles') or ('unconfirmedTitles' not in previous and (previous.get('errors') or 'after' not in previous)):
        raise ValueError('unresolved product title writes must be reconciled before OSW switch')
    original = path.read_bytes()
    archived = output / '.state/short-title-route-history' / (sha(original) + '.json')
    archived.parent.mkdir(parents=True, exist_ok=True)
    if archived.exists() and archived.read_bytes() != original:
        raise ValueError('original title route archive changed')
    archived.write_bytes(original)
    previous.update({'transport': 'osw', 'rejectedTitles': {}, 'previousRoute': {'path': str(archived), 'sha256': sha(original)}})
    save(path, previous)

def execute_task(client, body, output):
    rows = body.get('reqList', [])
    erp = getattr(client, 'expected_erp', '')
    if not erp or not rows or len({row['skuId'] for row in rows}) != len(rows) or any((set(row) != {'productId', 'skuId', 'shortTitle'} for row in rows)) or any((not str(row['skuId']).isdigit() or not str(row['productId']).isdigit() or (not isinstance(row['shortTitle'], str)) or (not row['shortTitle'].strip()) for row in rows)):
        raise ValueError('invalid authorized OSW short-title request')
    identity = {'erp': erp, 'rows': rows, 'type': 272}
    key = sha(json.dumps(identity, sort_keys=True, ensure_ascii=False).encode())
    folder = output / '.state/osw-title-tasks' / key
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / 'intent.json'
    if path.exists():
        record = read(path)
        if record['identity'] != identity:
            raise ValueError('OSW frozen task identity changed')
        if sha((folder / 'source.xlsx').read_bytes()) != record['sourceSha256']:
            raise ValueError('OSW frozen source file changed')
        if record.get('state') == 'complete':
            if sha((folder / 'downloaded-source.xlsx').read_bytes()) != record['sourceSha256'] or sha((folder / 'row-log.xlsx').read_bytes()) != record['logSha256']:
                raise ValueError('OSW cached completion files changed')
            return parse_log((folder / 'row-log.xlsx').read_bytes(), rows)
    else:
        data, _ = export.fill_template_bytes(export.TITLE_TEMPLATE, '批量维护短标题', [[row['skuId'], row['shortTitle']] for row in rows])
        if len(data) >= 5 * 1024 * 1024:
            raise ValueError('OSW frozen source exceeds native file limit')
        (folder / 'source.xlsx').write_bytes(data)
        name = 'CT短标题_' + key[:20]
        client.osw_title_action({'kind': 'prepare', 'erp': erp})
        record = {'identity': identity, 'name': name, 'sourceSha256': sha(data), 'state': 'create-unknown', 'startedAt': time.time()}
        save(path, record)
        reply = client.osw_title_action({'kind': 'create', 'erp': erp, 'name': name, 'sourceSha256': sha(data), 'base64': base64.b64encode(data).decode()})
        record['createReceipt'] = reply
        save(path, record)
        if reply.get('success') is not True or not re.fullmatch('[0-9a-f]{24}', str(reply.get('data', ''))):
            raise ValueError('OSW create response uncertain; reconcile task before any retry')
        record.update({'taskId': reply['data'], 'state': 'task-created'})
        save(path, record)
    deadline = time.monotonic() + 180
    while True:
        reply = client.osw_title_action({'kind': 'lookup', 'erp': erp, 'name': record['name']})
        task = checked_task(reply, record['name'], erp, record.get('taskId'))
        record.update({'taskId': task['id'], 'task': task, 'state': 'task-created', 'lastCheckedAt': time.time()})
        save(path, record)
        if task.get('status') in [0, 3]:
            break
        if time.monotonic() >= deadline:
            raise TimeoutError('OSW task still pending; retain original task and reconcile')
        time.sleep(2)
    for kind, filename in [('downsourcefile', 'downloaded-source.xlsx'), ('downlog', 'row-log.xlsx')]:
        reply = client.osw_title_action({'kind': 'file', 'erp': erp, 'name': record['name'], 'taskId': record['taskId'], 'fileKind': kind})
        data = base64.b64decode(reply['base64'], validate=True)
        if len(data) != reply['bytes'] or (kind == 'downsourcefile' and sha(data) != record['sourceSha256']):
            raise ValueError('OSW downloaded source or file size mismatch')
        (folder / filename).write_bytes(data)
    outcomes = parse_log((folder / 'row-log.xlsx').read_bytes(), rows)
    record.update({'state': 'complete', 'outcomes': outcomes, 'logSha256': sha((folder / 'row-log.xlsx').read_bytes()), 'elapsedSeconds': time.time() - record['startedAt'], 'completedAt': time.time()})
    save(path, record)
    return outcomes

def submit(client, spu_ids, desired, output, *, expected_before=None, defer_conflicts=False):
    import direct_binding as binding
    output = Path(output)
    audit_path = output / '.state/short-title-write.json'
    if audit_path.exists():
        previous = read(audit_path)
        if previous.get('transport') == 'osw' and previous.get('unconfirmedTitles'):
            for path in sorted((output / '.state/osw-title-tasks').glob('*/intent.json')):
                record = read(path)
                rows = record['identity']['rows']
                if not any((row['skuId'] in previous['unconfirmedTitles'] for row in rows)):
                    continue
                if any((desired.get(row['skuId']) != row['shortTitle'] for row in rows)):
                    raise ValueError('OSW uncertain original request differs from desired titles')
                outcomes = execute_task(client, {'reqList': rows}, output)
                by_sku = {item['skuId']: item for item in outcomes}
                for row in rows:
                    sku = row['skuId']
                    item = by_sku[sku]
                    previous['unconfirmedTitles'].pop(sku, None)
                    if item['success']:
                        previous.setdefault('submittedTitles', {})[sku] = row['shortTitle']
                    else:
                        previous.setdefault('rejectedTitles', {})[sku] = {'title': row['shortTitle'], 'reason': item['message'], 'code': item['code']}
                save(audit_path, previous)
            if previous.get('unconfirmedTitles'):
                raise ValueError('OSW uncertain titles have no resolved task receipt')
    prepare_route(output)

    class Proxy:
        short_title_transport = 'osw-active'
        title_audit_transport = 'osw'

        def __getattr__(self, name):
            return getattr(client, name)

        def save_short_titles(self, body):
            outcomes = execute_task(client, body, output)
            if any((not item['success'] for item in outcomes)):
                raise binding.ShortTitleBatchIncomplete('short-title submission incomplete; outcomes=' + json.dumps(outcomes, ensure_ascii=False))
            return {'ok': True, 'submitted': len(body['reqList'])}
    return binding.submit_self_operated_short_titles(Proxy(), spu_ids, desired, output, expected_before=expected_before, defer_conflicts=defer_conflicts)
